import io
import os
from datetime import datetime

import openpyxl
from flask import (
    Blueprint,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook
from sqlalchemy import func

from app.extensions import db
from app.models import ImportMapping, User
from app.services.dashboard_service import DashboardService
from app.services.email_service import EmailService
from app.services.import_service import ImportService
from app.services.order_service import OrderService
from app.services.user_service import UserService
from config import ADMIN_TOKEN, UPLOAD_FOLDER

admin_bp = Blueprint("admin", __name__, url_prefix="/admin")


def admin_required(f):
    from functools import wraps

    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get("is_admin"):
            return redirect(url_for("auth.admin_login"))
        return f(*args, **kwargs)

    return decorated_function


@admin_bp.route("/")
@admin_required
def dashboard():
    stats = DashboardService.admin_stats()

    upload_folder = os.path.abspath(UPLOAD_FOLDER)
    upload_size = 0
    if os.path.exists(upload_folder):
        for root, _dirs, files in os.walk(upload_folder):
            for f in files:
                fp = os.path.join(root, f)
                if os.path.isfile(fp):
                    upload_size += os.path.getsize(fp)

    stats["upload_folder"] = upload_folder
    stats["upload_size"] = upload_size

    return render_template("admin_dashboard.html", stats=stats)


@admin_bp.route("/pickers")
@admin_required
def pickers():
    pickers_data = DashboardService.picker_stats()
    return render_template("admin_pickers.html", pickers=pickers_data)


@admin_bp.route("/couriers")
@admin_required
def couriers():
    complaints_data = DashboardService.complaint_couriers()
    users_count = db.session.query(func.count(User.id)).scalar() or 0
    # Ждут фото: только подтемы «нужно фото» и фото ещё нет
    awaiting_photo = sum(
        1 for row in complaints_data if row["photo_required"] and not row["has_photos"]
    )
    return render_template(
        "admin_couriers.html",
        complaints=complaints_data,
        users_count=users_count,
        orders_with_complaints=len(complaints_data),
        awaiting_photo=awaiting_photo,
    )


@admin_bp.route("/api/pending_count")
def api_pending_count():
    """Живой счётчик заказов, ждущих подтверждения (для polling на дашборде)."""
    if not session.get("is_admin"):
        return jsonify({"error": "forbidden"}), 403
    return jsonify({"pending": DashboardService.pending_orders_count()})


@admin_bp.route("/api/files")
def api_files():
    """List uploaded files — used by the PC sync script.

    Auth: admin session OR X-Admin-Token header with the superadmin token.
    """
    token = request.headers.get("X-Admin-Token", "")
    if not session.get("is_admin") and token != ADMIN_TOKEN:
        return jsonify({"error": "unauthorized"}), 401

    folder = os.path.abspath(UPLOAD_FOLDER)
    files = []
    if os.path.isdir(folder):
        for name in sorted(os.listdir(folder)):
            fp = os.path.join(folder, name)
            if os.path.isfile(fp):
                files.append(
                    {
                        "name": name,
                        "size": os.path.getsize(fp),
                        "mtime": int(os.path.getmtime(fp)),
                    }
                )
    return jsonify({"files": files, "count": len(files)})


@admin_bp.route("/users", methods=["GET", "POST"])
@admin_required
def users():
    page = request.args.get("page", 1, type=int)
    search = request.args.get("search", "", type=str)

    if request.method == "POST":
        fio = request.form.get("fio", "").strip()
        if fio:
            try:
                user = UserService.create_user(fio)
                flash(f"Пользователь создан. Токен: {user.token}", "success")
            except Exception:
                db.session.rollback()
                flash("Пользователь с таким ФИО уже существует", "error")
        else:
            flash("Введите ФИО", "error")
        return redirect(url_for("admin.users"))

    pagination = UserService.get_all_users(page=page, search=search)
    return render_template(
        "admin_users.html", users=pagination.items, pagination=pagination, search=search
    )


@admin_bp.route("/users/delete/<int:user_id>", methods=["POST"])
@admin_required
def delete_user(user_id):
    user = UserService.delete_user(user_id)
    if user:
        flash("Пользователь удалён", "success")
    return redirect(url_for("admin.users"))


@admin_bp.route("/users/regenerate_token/<int:user_id>", methods=["POST"])
@admin_required
def regenerate_token(user_id):
    user = UserService.regenerate_token(user_id)
    if user:
        flash(f"Токен обновлён. Новый токен: {user.token}", "success")
    return redirect(url_for("admin.users"))


@admin_bp.route("/upload_excel", methods=["GET", "POST"])
@admin_required
def upload_excel():
    if request.method == "POST":
        file = request.files.get("excel_file")
        if not file or not file.filename:
            flash("Файл не выбран", "error")
            return redirect(url_for("admin.upload_excel"))

        if not file.filename.endswith((".xlsx", ".xls")):
            flash("Неверный формат файла. Нужен .xlsx или .xls", "error")
            return redirect(url_for("admin.upload_excel"))

        try:
            workbook = openpyxl.load_workbook(file)
            sheet = workbook.active

            headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]

            fio_col_idx = None
            fio_variants = ["фио", "fio", "ФИО", "Фио", "Фамилия Имя Отчество", "ФИО клиента"]
            for idx, header in enumerate(headers):
                if header and str(header).strip().lower() in [v.lower() for v in fio_variants]:
                    fio_col_idx = idx
                    break

            if fio_col_idx is None:
                flash(
                    'Не найдена колонка с ФИО. Ожидаемые названия: "ФИО", "fio", "ФИО клиента"',
                    "error",
                )
                return redirect(url_for("admin.upload_excel"))

            result = ImportService.import_orders_from_excel(file, fio_col_idx)

            msg = f'Импортировано заказов: {result["imported"]}. Пропущено строк: {result["skipped"]}. Дублей: {result["duplicates"]}.'
            if result["errors"]:
                msg += " Ошибки: " + "; ".join(result["errors"][:5])
                if len(result["errors"]) > 5:
                    msg += f' ... и ещё {len(result["errors"]) - 5} ошибок.'
            flash(msg, "success" if result["imported"] > 0 else "warning")

        except Exception as e:
            db.session.rollback()
            flash(f"Ошибка при обработке файла: {str(e)}", "error")

        return redirect(url_for("admin.upload_excel"))

    return render_template("admin_upload_excel.html")


@admin_bp.route("/import_mappings", methods=["GET", "POST"])
@admin_required
def import_mappings():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        type_ = request.form.get("type", "").strip()
        id_column = request.form.get("id_column", "").strip()
        fio_column = request.form.get("fio_column", "").strip()
        order_number_column = request.form.get("order_number_column", "").strip()

        field_mapping = {}
        keys = request.form.getlist("mapping_key[]")
        values = request.form.getlist("mapping_value[]")
        for k, v in zip(keys, values, strict=False):
            if k.strip() and v.strip():
                field_mapping[k.strip()] = v.strip()

        if not name or not type_ or not id_column:
            flash("Заполните обязательные поля: название, тип, колонка ID", "error")
        else:
            try:
                mapping = ImportMapping(
                    name=name,
                    type=type_,
                    id_column=id_column,
                    fio_column=fio_column or None,
                    order_number_column=order_number_column or None,
                    field_mapping=field_mapping,
                )
                db.session.add(mapping)
                db.session.commit()
                flash("Маппинг создан", "success")
            except Exception as e:
                db.session.rollback()
                flash(f"Ошибка: {e}", "error")
        return redirect(url_for("admin.import_mappings"))

    mappings = ImportMapping.query.order_by(ImportMapping.created_at.desc()).all()
    return render_template("admin_import_mappings.html", mappings=mappings)


@admin_bp.route("/import_mappings/delete/<int:mapping_id>", methods=["POST"])
@admin_required
def delete_import_mapping(mapping_id):
    mapping = ImportMapping.query.get_or_404(mapping_id)
    db.session.delete(mapping)
    db.session.commit()
    flash("Маппинг удалён", "success")
    return redirect(url_for("admin.import_mappings"))


@admin_bp.route("/import_excel", methods=["GET", "POST"])
@admin_required
def import_excel():
    mappings = ImportMapping.query.order_by(ImportMapping.type, ImportMapping.name).all()

    if request.method == "POST":
        mapping_id = request.form.get("mapping_id", type=int)
        file = request.files.get("excel_file")

        if not mapping_id:
            flash("Выберите маппинг", "error")
            return redirect(url_for("admin.import_excel"))

        if not file or not file.filename:
            flash("Файл не выбран", "error")
            return redirect(url_for("admin.import_excel"))

        if not file.filename.endswith((".xlsx", ".xls")):
            flash("Неверный формат файла. Нужен .xlsx или .xls", "error")
            return redirect(url_for("admin.import_excel"))

        mapping = ImportMapping.query.get_or_404(mapping_id)

        try:
            result = ImportService.import_with_mapping(file, mapping)

            msg = f'Импорт завершён. Создано: {result["imported"]}, обновлено: {result["updated"]}, пропущено: {result["skipped"]}.'
            if result["errors"]:
                msg += " Ошибки: " + "; ".join(result["errors"][:5])
                if len(result["errors"]) > 5:
                    msg += f' ... и ещё {len(result["errors"]) - 5} ошибок.'
            flash(msg, "success" if (result["imported"] + result["updated"]) > 0 else "warning")

        except Exception as e:
            db.session.rollback()
            flash(f"Ошибка при обработке файла: {str(e)}", "error")

        return redirect(url_for("admin.import_excel"))

    return render_template("admin_import_excel.html", mappings=mappings)


@admin_bp.route("/orders")
@admin_required
def orders():
    page = request.args.get("page", 1, type=int)
    search = request.args.get("search", "", type=str)
    user_filter = request.args.get("user_id", "", type=str)
    photo_filter = request.args.get("photo_filter", "", type=str)

    pagination, orders_with_users = OrderService.get_all_orders(
        page=page, search=search, user_filter=user_filter, photo_filter=photo_filter
    )

    users = UserService.get_all_users(per_page=1000).items

    return render_template(
        "admin_orders.html",
        orders=orders_with_users,
        pagination=pagination,
        search=search,
        users=users,
        user_filter=user_filter,
        photo_filter=photo_filter,
    )


@admin_bp.route("/orders/send_email/<int:order_id>", methods=["POST"])
@admin_required
def send_order_email(order_id):
    result = OrderService.get_order_with_user(order_id)
    if not result:
        flash("Заказ не найден", "error")
        return redirect(url_for("admin.orders"))

    order, user = result

    if EmailService.send_order_completion_email(order, user):
        flash("Email отправлен администраторам", "success")
    else:
        flash("Ошибка отправки (проверьте настройки почты)", "error")

    return redirect(url_for("admin.orders"))


@admin_bp.route("/orders/delete/<int:order_id>", methods=["POST"])
@admin_required
def delete_order(order_id):
    if OrderService.delete_order(order_id):
        flash("Заказ удалён", "success")
    return redirect(url_for("admin.orders"))


@admin_bp.route("/orders/edit/<int:order_id>", methods=["GET", "POST"])
@admin_required
def edit_order(order_id):
    order = OrderService.get_order_by_id(order_id)
    if not order:
        flash("Заказ не найден", "error")
        return redirect(url_for("admin.orders"))

    if request.method == "POST":
        OrderService.update_order(
            order_id, request.form.get("order_number", ""), request.form.get("description", "")
        )
        flash("Заказ обновлён", "success")
        return redirect(url_for("admin.orders"))

    return render_template("admin_edit_order.html", order=order)


@admin_bp.route("/export_excel")
@admin_required
def export_excel():
    from app.models import Order, User

    wb = Workbook()
    ws = wb.active
    ws.title = "Заказы"

    headers = ["ID", "ФИО", "Токен", "Номер заказа", "Описание", "Создан", "Кол-во фото"]
    ws.append(headers)

    orders = db.session.query(Order, User).join(User).order_by(Order.created_at.desc()).all()

    for order, user in orders:
        ws.append(
            [
                order.id,
                user.fio,
                user.token,
                order.order_number,
                order.description,
                order.created_at.strftime("%Y-%m-%d %H:%M:%S") if order.created_at else "",
                len(order.photos),
            ]
        )

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f'orders_export_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx',
    )
