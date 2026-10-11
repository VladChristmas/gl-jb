from datetime import datetime

import openpyxl
from werkzeug.datastructures import FileStorage

from app.constants import PHOTO_REQUIRED_TOPICS
from app.extensions import db
from app.models.import_mapping import ImportMapping
from app.models.order import Order
from app.models.user import User
from app.services.audit_service import AuditService
from app.services.complaint_service import ComplaintService
from app.services.user_service import UserService


class ImportService:
    # Placeholder user for complaints whose order is missing from the orders import
    PLACEHOLDER_USER_FIO = "Не указан"
    # Excel header keyword -> Order column (first match wins, case-insensitive)
    DASHBOARD_FIELD_KEYWORDS = (
        ("курьер", "courier_fio"),
        ("адрес", "address"),
        ("доставлен", "delivered_at"),
        ("сборщик", "picker_fio"),
        ("собран", "pick_count"),
        ("ожидан", "wait_time"),
        ("скорость", "pick_speed"),
    )

    @staticmethod
    def _extract_dashboard_fields(headers: list, row: tuple) -> dict:
        """Extract courier/picker columns from an Excel row by header keywords."""
        found: dict = {}
        for idx, header in enumerate(headers):
            if header is None or idx >= len(row):
                continue
            text = str(header).strip().lower()
            if not text:
                continue
            for keyword, field in ImportService.DASHBOARD_FIELD_KEYWORDS:
                if keyword in text and field not in found:
                    value = row[idx]
                    if value is None or str(value).strip() == "":
                        break
                    if field == "pick_count":
                        try:
                            found[field] = int(float(str(value).strip().replace(",", ".")))
                        except ValueError:
                            pass
                    else:
                        found[field] = str(value).strip()
                    break
        return found

    @staticmethod
    def _number_conflict(user_id: int, order_number: str, exclude_id: int | None = None) -> bool:
        """Занят ли (user_id, order_number) другим заказом.

        Общий guard для обоих путей импорта (быстрый путь проверял это перед
        INSERT): без него нарушение уникального индекса откатывало весь файл.
        """
        query = Order.query.filter_by(user_id=user_id, order_number=order_number)
        if exclude_id is not None:
            query = query.filter(Order.id != exclude_id)
        return query.first() is not None

    @staticmethod
    def import_orders_from_excel(file: FileStorage, fio_col_idx: int) -> dict:
        workbook = openpyxl.load_workbook(file)
        sheet = workbook.active

        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        users_map = UserService.get_users_map()

        imported = 0
        skipped = 0
        duplicates = 0
        errors = []

        for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(row):
                continue

            fio_value = row[fio_col_idx]
            if not fio_value:
                skipped += 1
                errors.append(f"Строка {row_idx}: пустое ФИО")
                continue

            fio_normalized = str(fio_value).strip().lower()
            user_id = users_map.get(fio_normalized)

            if not user_id:
                skipped += 1
                errors.append(f'Строка {row_idx}: пользователь "{fio_value}" не найден')
                continue

            order_data = {}
            for idx, header in enumerate(headers):
                if idx != fio_col_idx and header:
                    order_data[str(header).strip()] = row[idx]

            order_number = (
                order_data.get("Номер заказа")
                or order_data.get("order_number")
                or f"Заказ #{row_idx}"
            )

            existing = Order.query.filter_by(user_id=user_id, order_number=order_number).first()
            if existing:
                duplicates += 1
                errors.append(
                    f'Строка {row_idx}: заказ "{order_number}" уже существует для этого пользователя'
                )
                continue
            description_parts = []
            for key, value in order_data.items():
                if value is not None:
                    description_parts.append(f"{key}: {value}")
            description = "; ".join(description_parts)

            dash_fields = ImportService._extract_dashboard_fields(headers, row)

            order = Order(
                user_id=user_id,
                order_number=order_number,
                description=description,
                **dash_fields,
            )
            db.session.add(order)
            imported += 1

        db.session.commit()

        return {
            "imported": imported,
            "skipped": skipped,
            "duplicates": duplicates,
            "errors": errors,
        }

    @staticmethod
    def import_with_mapping(file: FileStorage, mapping: ImportMapping) -> dict:
        workbook = openpyxl.load_workbook(file)
        sheet = workbook.active

        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        headers = [str(h).strip() if h else "" for h in headers]

        id_col_idx = None
        for idx, header in enumerate(headers):
            if header.lower() == mapping.id_column.lower():
                id_col_idx = idx
                break

        if id_col_idx is None:
            raise ValueError(
                f'Колонка ID "{mapping.id_column}" не найдена в файле. Доступные: {headers}'
            )

        fio_col_idx = None
        if mapping.fio_column:
            for idx, header in enumerate(headers):
                if header.lower() == mapping.fio_column.lower():
                    fio_col_idx = idx
                    break

        order_num_col_idx = None
        if mapping.order_number_column:
            for idx, header in enumerate(headers):
                if header.lower() == mapping.order_number_column.lower():
                    order_num_col_idx = idx
                    break

        field_indices = {}
        for file_col, model_field in mapping.field_mapping.items():
            for idx, header in enumerate(headers):
                if header.lower() == file_col.lower():
                    field_indices[model_field] = idx
                    break

        users_map = UserService.get_users_map()

        imported = 0
        updated = 0
        skipped = 0
        duplicates = 0
        errors = []
        # §8: строки с ID, уже встречавшимся в этом файле, считаются дублями
        # (повторная строка перезаписывает первую — поведение сохранено)
        seen_ids: set[str] = set()

        start_row = 2 if mapping.skip_first_row else 1

        for row_idx, row in enumerate(
            sheet.iter_rows(min_row=start_row, values_only=True), start=start_row
        ):
            if not any(row):
                continue

            ext_id = row[id_col_idx]
            if not ext_id:
                skipped += 1
                errors.append(f"Строка {row_idx}: пустой ID")
                continue

            ext_id = str(ext_id).strip()
            if ext_id in seen_ids:
                duplicates += 1
            seen_ids.add(ext_id)

            if mapping.type == "orders":
                fio_value = row[fio_col_idx] if fio_col_idx is not None else None
                if not fio_value:
                    skipped += 1
                    errors.append(f"Строка {row_idx}: пустое ФИО")
                    continue

                fio_normalized = str(fio_value).strip().lower()
                user_id = users_map.get(fio_normalized)
                if not user_id:
                    skipped += 1
                    errors.append(f'Строка {row_idx}: пользователь "{fio_value}" не найден')
                    continue

                order_number = (
                    row[order_num_col_idx] if order_num_col_idx is not None else f"Заказ #{row_idx}"
                )
                if not order_number:
                    order_number = f"Заказ #{row_idx}"

                order = Order.query.filter_by(external_id=ext_id).first()

                description_parts = []
                for model_field, col_idx in field_indices.items():
                    value = row[col_idx]
                    if value is not None:
                        description_parts.append(f"{model_field}: {value}")
                description = "; ".join(description_parts)

                dash_fields = ImportService._extract_dashboard_fields(headers, row)

                if order:
                    # Конфликт уникальности (user_id, order_number): чужой заказ
                    # с тем же номером — не переписываем его, считаем дублем.
                    # Раньше это падало IntegrityError'ом и откатывало весь импорт.
                    conflict = ImportService._number_conflict(
                        user_id, order_number, exclude_id=order.id
                    )
                    if conflict:
                        duplicates += 1
                        errors.append(
                            f'Строка {row_idx}: номер "{order_number}" уже занят '
                            f"другим заказом этого пользователя"
                        )
                        continue
                    order.user_id = user_id
                    order.order_number = order_number
                    order.description = description
                    order.external_id = ext_id
                    for field, value in dash_fields.items():
                        setattr(order, field, value)
                    updated += 1
                elif ImportService._number_conflict(user_id, order_number):
                    duplicates += 1
                    errors.append(
                        f'Строка {row_idx}: заказ "{order_number}" уже существует '
                        f"для этого пользователя (другой ID)"
                    )
                    continue
                else:
                    order = Order(
                        user_id=user_id,
                        order_number=order_number,
                        description=description,
                        external_id=ext_id,
                        **dash_fields,
                    )
                    db.session.add(order)
                    imported += 1

            elif mapping.type == "complaints":
                order = Order.query.filter_by(external_id=ext_id).first()
                if not order:
                    # Complaint ID may hold the order number instead of external ID
                    order = Order.query.filter_by(order_number=ext_id).first()

                # Значение подтемы из файла (для определения, нужно ли фото)
                topic_value = ""
                for model_field, col_idx in field_indices.items():
                    if model_field == "complaint_text" and row[col_idx] is not None:
                        topic_value = str(row[col_idx]).strip()

                if not order:
                    needs_photo = any(t in topic_value for t in PHOTO_REQUIRED_TOPICS)
                    if not needs_photo:
                        # Заказ не найден и подтема не требует фото — пропускаем
                        skipped += 1
                        errors.append(f'Строка {row_idx}: заказ с ID "{ext_id}" не найден')
                        continue
                    # Заказ не найден, но подтема требует фото — заказ-заглушка,
                    # чтобы жалоба с фото-подтемой не потерялась
                    user = User.query.filter_by(fio=ImportService.PLACEHOLDER_USER_FIO).first()
                    if not user:
                        user = UserService.create_user(ImportService.PLACEHOLDER_USER_FIO)
                    order = Order(
                        user_id=user.id,
                        external_id=ext_id,
                        order_number=f"Жалоба {ext_id[:8]}",
                    )
                    db.session.add(order)
                    ComplaintService.add(order, topic_value)
                    imported += 1
                    continue

                # Одна строка Excel = одна жалоба (текст + статус + дата).
                # Раньше status/date перезаписывались каждой строкой и терялись.
                new_text = None
                new_status = None
                new_date = None
                for model_field, col_idx in field_indices.items():
                    value = row[col_idx]
                    if value is None:
                        continue
                    if model_field == "complaint_text":
                        new_text = str(value).strip()
                    elif model_field == "complaint_status":
                        new_status = str(value)
                    elif model_field == "complaint_date":
                        if isinstance(value, str):
                            try:
                                new_date = datetime.fromisoformat(value.replace("Z", "+00:00"))
                            except Exception:
                                new_date = None
                        elif hasattr(value, "isoformat"):
                            new_date = value

                if new_text:
                    ComplaintService.add(order, new_text, new_status, new_date)
                elif new_status or new_date:
                    # Строка без текста — правим статус/дату последней жалобы
                    ComplaintService.hydrate_from_mirror(order)
                    if order.complaints:
                        last = order.complaints[-1]
                        if new_status:
                            last.status = new_status
                        if new_date:
                            last.created_at = new_date
                        ComplaintService.sync_mirror(order)
                    else:
                        if new_status:
                            order.complaint_status = new_status
                        if new_date:
                            order.complaint_date = new_date

                updated += 1

        db.session.commit()

        # Audit log for import (single entry with summary)
        result_dict = {
            "imported": imported,
            "updated": updated,
            "skipped": skipped,
            "duplicates": duplicates,
            "errors": errors,
        }
        AuditService.log_import(mapping.name, mapping.type, result_dict)

        return result_dict

    @staticmethod
    def preview_import(file: FileStorage, mapping: ImportMapping) -> dict:
        """
        Предварительный анализ файла без сохранения в БД.
        Возвращает структурированный превью с категориями строк.
        """
        file.seek(0)
        workbook = openpyxl.load_workbook(file)
        sheet = workbook.active

        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        headers = [str(h).strip() if h else "" for h in headers]

        id_col_idx = None
        for idx, header in enumerate(headers):
            if header.lower() == mapping.id_column.lower():
                id_col_idx = idx
                break

        if id_col_idx is None:
            raise ValueError(
                f'Колонка ID "{mapping.id_column}" не найдена в файле. Доступные: {headers}'
            )

        fio_col_idx = None
        if mapping.fio_column:
            for idx, header in enumerate(headers):
                if header.lower() == mapping.fio_column.lower():
                    fio_col_idx = idx
                    break

        order_num_col_idx = None
        if mapping.order_number_column:
            for idx, header in enumerate(headers):
                if header.lower() == mapping.order_number_column.lower():
                    order_num_col_idx = idx
                    break

        field_indices = {}
        for file_col, model_field in mapping.field_mapping.items():
            for idx, header in enumerate(headers):
                if header.lower() == file_col.lower():
                    field_indices[model_field] = idx
                    break

        users_map = UserService.get_users_map()

        # Результаты превью
        new_rows = []  # новые заказы/жалобы
        updated_rows = []  # обновляемые существующие
        duplicate_rows = []  # дубликаты в файле
        conflict_rows = []  # конфликты с существующими данными
        unmatchable_rows = []  # не сопоставленные (нет заказа, нет юзера)
        error_rows = []  # ошибки формата/валидации
        skipped_rows = []  # пропущенные (пустые ID, пустое ФИО)

        # Для отслеживания дублей внутри файла
        seen_ids: set[str] = set()
        seen_order_numbers: dict[str, int] = {}  # (user_id, order_number) -> row_idx

        start_row = 2 if mapping.skip_first_row else 1

        for row_idx, row in enumerate(
            sheet.iter_rows(min_row=start_row, values_only=True), start=start_row
        ):
            if not any(row):
                continue

            ext_id = row[id_col_idx]
            if not ext_id:
                skipped_rows.append(
                    {
                        "row": row_idx,
                        "values": [str(v) if v is not None else "" for v in row],
                        "reason": "Пустой ID",
                    }
                )
                continue

            ext_id = str(ext_id).strip()

            # Проверка дублей внутри файла
            if ext_id in seen_ids:
                duplicate_rows.append(
                    {
                        "row": row_idx,
                        "values": [str(v) if v is not None else "" for v in row],
                        "external_id": ext_id,
                        "reason": f"Дубликат ID в файле: {ext_id}",
                    }
                )
                continue
            seen_ids.add(ext_id)

            # Собираем распознанные значения для превью
            recognized = {}
            for model_field, col_idx in field_indices.items():
                value = row[col_idx]
                if value is not None:
                    recognized[model_field] = str(value)

            if mapping.type == "orders":
                fio_value = row[fio_col_idx] if fio_col_idx is not None else None
                if not fio_value:
                    skipped_rows.append(
                        {
                            "row": row_idx,
                            "values": [str(v) if v is not None else "" for v in row],
                            "reason": "Пустое ФИО",
                        }
                    )
                    continue

                fio_normalized = str(fio_value).strip().lower()
                user_id = users_map.get(fio_normalized)
                if not user_id:
                    unmatchable_rows.append(
                        {
                            "row": row_idx,
                            "values": [str(v) if v is not None else "" for v in row],
                            "recognized": recognized,
                            "reason": f'Пользователь "{fio_value}" не найден в системе',
                        }
                    )
                    continue

                order_number = (
                    row[order_num_col_idx] if order_num_col_idx is not None else f"Заказ #{row_idx}"
                )
                if not order_number:
                    order_number = f"Заказ #{row_idx}"

                order_number = str(order_number).strip()

                # Ключ уникальности: (user_id, order_number)
                order_key = f"{user_id}:{order_number}"

                # Проверяем существующий заказ по external_id
                existing_by_ext = Order.query.filter_by(external_id=ext_id).first()

                if existing_by_ext:
                    # Проверяем конфликт номера заказа с чужим заказом
                    conflict = ImportService._number_conflict(
                        user_id, order_number, exclude_id=existing_by_ext.id
                    )
                    if conflict:
                        conflict_rows.append(
                            {
                                "row": row_idx,
                                "values": [str(v) if v is not None else "" for v in row],
                                "recognized": recognized,
                                "reason": f'Номер "{order_number}" уже занят другим заказом этого пользователя',
                                "existing_id": existing_by_ext.id,
                            }
                        )
                        continue

                    # Будет обновлён
                    updated_rows.append(
                        {
                            "row": row_idx,
                            "values": [str(v) if v is not None else "" for v in row],
                            "recognized": recognized,
                            "action": "update",
                            "external_id": ext_id,
                            "existing_id": existing_by_ext.id,
                            "existing_external_id": existing_by_ext.external_id,
                            "existing_order_number": existing_by_ext.order_number,
                            "new_order_number": order_number,
                            "user_id": user_id,
                            "fio": fio_value,
                        }
                    )
                elif ImportService._number_conflict(user_id, order_number):
                    # Конфликт с другим заказом (не по external_id)
                    conflict_rows.append(
                        {
                            "row": row_idx,
                            "values": [str(v) if v is not None else "" for v in row],
                            "recognized": recognized,
                            "external_id": ext_id,
                            "reason": f'Заказ "{order_number}" уже существует для этого пользователя (другой ID)',
                            "order_number": order_number,
                            "user_id": user_id,
                        }
                    )
                    continue
                else:
                    # Новый заказ
                    # Проверяем дубликаты order_number в этом же импорте
                    if order_key in seen_order_numbers:
                        conflict_rows.append(
                            {
                                "row": row_idx,
                                "values": [str(v) if v is not None else "" for v in row],
                                "recognized": recognized,
                                "external_id": ext_id,
                                "reason": f'Номер "{order_number}" уже встречается в этом импорте (строка {seen_order_numbers[order_key]})',
                            }
                        )
                        continue
                    seen_order_numbers[order_key] = row_idx

                    new_rows.append(
                        {
                            "row": row_idx,
                            "values": [str(v) if v is not None else "" for v in row],
                            "recognized": recognized,
                            "action": "create",
                            "order_number": order_number,
                            "user_id": user_id,
                            "fio": fio_value,
                            "external_id": ext_id,
                        }
                    )

            elif mapping.type == "complaints":
                # Для жалоб ищем заказ по external_id или order_number
                order = Order.query.filter_by(external_id=ext_id).first()
                if not order:
                    order = Order.query.filter_by(order_number=ext_id).first()

                topic_value = ""
                for model_field, col_idx in field_indices.items():
                    if model_field == "complaint_text" and row[col_idx] is not None:
                        topic_value = str(row[col_idx]).strip()

                needs_photo = any(t in topic_value for t in PHOTO_REQUIRED_TOPICS)

                if not order:
                    if not needs_photo:
                        unmatchable_rows.append(
                            {
                                "row": row_idx,
                                "values": [str(v) if v is not None else "" for v in row],
                                "recognized": recognized,
                                "reason": f'Заказ с ID "{ext_id}" не найден, подтема не требует фото',
                            }
                        )
                        continue
                    # Заказ не найден, но нужна фото-подтема — будет создан заказ-заглушка
                    new_rows.append(
                        {
                            "row": row_idx,
                            "values": [str(v) if v is not None else "" for v in row],
                            "recognized": recognized,
                            "action": "create_placeholder",
                            "reason": "Заказ не найден, будет создан заказ-заглушка для фото-подтемы",
                            "external_id": ext_id,
                            "topic": topic_value,
                        }
                    )
                else:
                    # Жалоба будет добавлена/обновлена к существующему заказу
                    updated_rows.append(
                        {
                            "row": row_idx,
                            "values": [str(v) if v is not None else "" for v in row],
                            "recognized": recognized,
                            "action": "add_complaint",
                            "existing_id": order.id,
                            "existing_external_id": order.external_id,
                            "existing_order_number": order.order_number,
                            "topic": topic_value,
                        }
                    )

        return {
            "total_rows": len(seen_ids) + len(duplicate_rows) + len(skipped_rows),
            "new": {
                "count": len(new_rows),
                "rows": new_rows,
            },
            "updated": {
                "count": len(updated_rows),
                "rows": updated_rows,
            },
            "duplicates": {
                "count": len(duplicate_rows),
                "rows": duplicate_rows,
            },
            "conflicts": {
                "count": len(conflict_rows),
                "rows": conflict_rows,
            },
            "unmatchable": {
                "count": len(unmatchable_rows),
                "rows": unmatchable_rows,
            },
            "skipped": {
                "count": len(skipped_rows),
                "rows": skipped_rows,
            },
            "errors": {
                "count": len(error_rows),
                "rows": error_rows,
            },
            "has_data": len(new_rows) > 0 or len(updated_rows) > 0,
        }
