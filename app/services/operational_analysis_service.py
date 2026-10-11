"""Сервис анализа операционных данных даркстора.

Обрабатывает Excel-файл с тремя листами (жалобы, заказы, собранные заказы),
связывает их по ID заказа, рассчитывает метрики сборщиков и жалобы.
"""

import io
import re
import statistics
from collections import defaultdict
from datetime import date, datetime, time

import openpyxl
from openpyxl import Workbook

from app.extensions import db
from app.models.operational_analysis import (
    AssemblyOrder,
    DataQualityIssue,
    OperationalAnalysisSession,
    PickerMetrics,
)


class ColumnDetectionResult:
    """Результат автоопределения столбцов на листе."""

    def __init__(self):
        self.columns: dict[str, int] = {}
        self.unused: list[str] = []
        self.date_format_hints: list[str] = []


class OperationalAnalysisService:
    """Обработка Excel с 3 листами и расчёт операционных метрик."""

    DEFAULT_WAIT_THRESHOLD = 30

    # Варианты названий столбцов для каждого листа
    COMPLAINTS_SHEET_NAMES = ["жалобы", "лист1", "complaints", "л1"]
    ORDERS_SHEET_NAMES = ["заказы", "лист2", "orders", "л2"]
    ASSEMBLY_SHEET_NAMES = ["собранные заказы", "лист3", "assembly", "собрано", "л3"]

    # Ключевые слова для определения столбцов
    ORDER_ID_KEYWORDS = ["номер заказа", "id заказа", "заказ", "order id", "order_id", "номер"]
    PICKER_KEYWORDS = ["сборщик", "picker", "собрал", "кто собрал"]
    START_TIME_KEYWORDS = [
        "начало сборки",
        "время начала",
        "начало",
        "start",
        "время старта",
        "собрано с",
        "начал собирать",
    ]
    END_TIME_KEYWORDS = [
        "завершение сборки",
        "конец сборки",
        "время окончания",
        "окончание",
        "конец",
        "end",
        "собрано в",
        "собрано до",
        "закончил собирать",
    ]
    QUEUE_TIME_KEYWORDS = [
        "поступление в очередь",
        "время поступления",
        "очередь",
        "принят",
        "принят в работу",
        "дата заказа",
        "создан",
        "created",
    ]
    COURIER_KEYWORDS = ["курьер", "delivery", "доставлял"]
    ADDRESS_KEYWORDS = ["адрес", "address"]
    ITEMS_KEYWORDS = [
        "количество товаров",
        "кол-во товаров",
        "товаров",
        "items",
        "кол-во позиций",
        "позиций",
        "штук",
        "кол-во",
    ]
    COMPLAINT_TEXT_KEYWORDS = ["жалоба", "причина", "текст жалобы", "подтема", "тема", "complaint"]

    @staticmethod
    def _normalize_id(value) -> str:
        """Нормализация ID заказа: убирает пробелы, сохраняет значащие нули."""
        if value is None:
            return ""
        if isinstance(value, float):
            if value == int(value):
                return str(int(value))
            return str(value)
        if isinstance(value, int):
            return str(value)
        text = str(value).strip()
        # Убираем .0 в конце если это целое
        if re.match(r"^\d+\.0$", text):
            return text[:-2]
        return text

    @staticmethod
    def _parse_datetime(value):
        """Парсинг даты/времени из Excel (datetime, строка, дробь)."""
        if value is None:
            return None
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime.combine(value, time.min)
        if isinstance(value, int | float):
            # Excel serial date
            try:
                if value > 20000:  # > 1954 года
                    from openpyxl.utils.datetime import from_excel

                    return from_excel(value)
            except Exception:
                pass
            return None
        text = str(value).strip()
        if not text:
            return None
        # Форматы дат
        formats = [
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d %H:%M",
            "%d.%m.%Y %H:%M:%S",
            "%d.%m.%Y %H:%M",
            "%d/%m/%Y %H:%M:%S",
            "%d/%m/%Y %H:%M",
            "%Y-%m-%d",
            "%d.%m.%Y",
            "%d/%m/%Y",
            "%H:%M:%S",
            "%H:%M",
        ]
        for fmt in formats:
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                continue
        # ISO format
        try:
            return datetime.fromisoformat(text.replace("Z", "+00:00"))
        except Exception:
            return None

    @staticmethod
    def _parse_number(value) -> float | None:
        """Парсинг числа из строки ('5', '5,5', '5.5', '5 мин')."""
        if value is None:
            return None
        if isinstance(value, int | float):
            return float(value)
        text = str(value).strip()
        match = re.search(r"-?\d+(?:[.,]\d+)?", text)
        if match:
            return float(match.group(0).replace(",", "."))
        return None

    @classmethod
    def _detect_columns(cls, headers: list, keywords: dict[str, list[str]]) -> dict[str, int]:
        """Определение индексов столбцов по ключевым словам."""
        result: dict[str, int] = {}
        used: set[int] = set()
        normalized_headers = [str(h).strip().lower() if h else "" for h in headers]

        for field, field_keywords in keywords.items():
            if field in result:
                continue
            for idx, header in enumerate(normalized_headers):
                if idx in used or not header:
                    continue
                for kw in field_keywords:
                    if kw in header:
                        result[field] = idx
                        used.add(idx)
                        break
                if field in result:
                    break
        return result

    @classmethod
    def _detect_sheet(cls, workbook, candidates: list[str]):
        """Определение листа по названию."""
        sheet_names = [name.strip().lower() for name in workbook.sheetnames]
        for candidate in candidates:
            for idx, name in enumerate(sheet_names):
                if candidate in name:
                    return workbook.worksheets[idx]
        return None

    @classmethod
    def analyze_file(cls, file, wait_threshold: int | None = None) -> OperationalAnalysisSession:
        """Полный анализ Excel-файла: чтение, связывание, расчёт, сохранение.

        Принимает FileStorage или BytesIO.
        """
        if wait_threshold is None:
            wait_threshold = cls.DEFAULT_WAIT_THRESHOLD

        filename = getattr(file, "filename", None) or "unknown.xlsx"

        file.seek(0)
        workbook = openpyxl.load_workbook(file, data_only=True)

        session = OperationalAnalysisSession(
            filename=filename,
            status="processing",
            wait_threshold_minutes=wait_threshold,
        )
        db.session.add(session)
        db.session.flush()

        issues: list[DataQualityIssue] = []

        # Определяем листы
        complaints_sheet = cls._detect_sheet(workbook, cls.COMPLAINTS_SHEET_NAMES)
        orders_sheet = cls._detect_sheet(workbook, cls.ORDERS_SHEET_NAMES)
        assembly_sheet = cls._detect_sheet(workbook, cls.ASSEMBLY_SHEET_NAMES)

        # Если листы не распознаны по названию — пробуем по структуре
        if not complaints_sheet or not orders_sheet or not assembly_sheet:
            complaints_sheet, orders_sheet, assembly_sheet = cls._detect_sheets_by_structure(
                workbook, complaints_sheet, orders_sheet, assembly_sheet
            )

        # Читаем листы
        complaints_data, complaints_issue = cls._read_complaints_sheet(complaints_sheet)
        if complaints_issue:
            issues.append(complaints_issue)

        orders_data, orders_issue = cls._read_orders_sheet(orders_sheet)
        if orders_issue:
            issues.append(orders_issue)

        assembly_data, assembly_issue = cls._read_assembly_sheet(assembly_sheet)
        if assembly_issue:
            issues.append(assembly_issue)

        session.total_complaints = len(complaints_data)
        session.total_orders = len(orders_data)
        session.total_assembly_orders = len(assembly_data)

        # Строим индексы
        orders_by_id: dict[str, dict] = {}
        for order in orders_data:
            oid = order["order_id"]
            if oid in orders_by_id:
                issues.append(
                    DataQualityIssue(
                        session_id=session.id,
                        issue_type="duplicate_order",
                        description=f'Дубликат ID заказа в Листе2: "{oid}"',
                        source_sheet="orders",
                        order_id=oid,
                        severity="warning",
                    )
                )
            else:
                orders_by_id[oid] = order

        complaints_by_order: dict[str, list[dict]] = defaultdict(list)
        for complaint in complaints_data:
            oid = complaint["order_id"]
            complaints_by_order[oid].append(complaint)

        assembly_by_order: dict[str, dict] = {}
        for asm in assembly_data:
            oid = asm["order_id"]
            if oid in assembly_by_order:
                issues.append(
                    DataQualityIssue(
                        session_id=session.id,
                        issue_type="duplicate_assembly",
                        description=f'Дубликат ID заказа в Листе3: "{oid}"',
                        source_sheet="assembly",
                        order_id=oid,
                        severity="warning",
                    )
                )
            else:
                assembly_by_order[oid] = asm

        # Связываем данные
        matched_complaints = 0
        unmatched_complaints = 0
        for oid, complaints in complaints_by_order.items():
            if oid in orders_by_id:
                matched_complaints += len(complaints)
                orders_by_id[oid]["complaints"] = complaints
                orders_by_id[oid]["has_complaint"] = True
            else:
                unmatched_complaints += len(complaints)
                for _c in complaints:
                    issues.append(
                        DataQualityIssue(
                            session_id=session.id,
                            issue_type="unmatched_complaint",
                            description=f'Жалоба без соответствующего заказа: ID "{oid}"',
                            source_sheet="complaints",
                            order_id=oid,
                            severity="warning",
                        )
                    )

        matched_assembly = 0
        unmatched_assembly = 0
        for oid, asm in assembly_by_order.items():
            if oid in orders_by_id:
                matched_assembly += 1
                orders_by_id[oid]["assembly"] = asm
            else:
                unmatched_assembly += 1
                issues.append(
                    DataQualityIssue(
                        session_id=session.id,
                        issue_type="unmatched_assembly",
                        description=f'Собранный заказ без записи в Листе2: ID "{oid}"',
                        source_sheet="assembly",
                        order_id=oid,
                        severity="warning",
                    )
                )

        session.matched_complaints = matched_complaints
        session.matched_assembly = matched_assembly
        session.matched_orders = sum(1 for o in orders_by_id.values() if o.get("assembly"))

        # Сохраняем AssemblyOrder и считаем метрики
        picker_data: dict[str, list[dict]] = defaultdict(list)

        for oid, order in orders_by_id.items():
            assembly = order.get("assembly")
            complaints = order.get("complaints", [])
            complaint_count = len(complaints)
            has_complaint = complaint_count > 0

            if assembly:
                start = assembly.get("assembly_start")
                end = assembly.get("assembly_end")
                duration = None
                wait_minutes = assembly.get("wait_time_minutes")

                if start and end:
                    delta = end - start
                    duration = delta.total_seconds() / 60.0
                    if duration < 0:
                        issues.append(
                            DataQualityIssue(
                                session_id=session.id,
                                issue_type="negative_interval",
                                description=f'Отрицательное время сборки: ID "{oid}", начало {start}, конец {end}',
                                source_sheet="assembly",
                                order_id=oid,
                                picker_fio=assembly.get("picker_fio", ""),
                                severity="error",
                            )
                        )
                        duration = None
                    elif duration == 0:
                        issues.append(
                            DataQualityIssue(
                                session_id=session.id,
                                issue_type="zero_interval",
                                description=f'Нулевое время сборки: ID "{oid}"',
                                source_sheet="assembly",
                                order_id=oid,
                                picker_fio=assembly.get("picker_fio", ""),
                                severity="warning",
                            )
                        )

                if not start:
                    issues.append(
                        DataQualityIssue(
                            session_id=session.id,
                            issue_type="missing_timestamp",
                            description=f'Отсутствует время начала сборки: ID "{oid}"',
                            source_sheet="assembly",
                            order_id=oid,
                            picker_fio=assembly.get("picker_fio", ""),
                            severity="warning",
                        )
                    )
                if not end:
                    issues.append(
                        DataQualityIssue(
                            session_id=session.id,
                            issue_type="missing_timestamp",
                            description=f'Отсутствует время окончания сборки: ID "{oid}"',
                            source_sheet="assembly",
                            order_id=oid,
                            picker_fio=assembly.get("picker_fio", ""),
                            severity="warning",
                        )
                    )

                # Ожидание: от поступления в очередь до начала сборки
                if wait_minutes is None and order.get("queue_time") and start:
                    wait_delta = start - order["queue_time"]
                    wait_minutes = wait_delta.total_seconds() / 60.0
                    if wait_minutes < 0:
                        issues.append(
                            DataQualityIssue(
                                session_id=session.id,
                                issue_type="negative_wait",
                                description=f'Отрицательное ожидание: ID "{oid}"',
                                source_sheet="assembly",
                                order_id=oid,
                                severity="warning",
                            )
                        )
                        wait_minutes = None

                ao = AssemblyOrder(
                    session_id=session.id,
                    order_id=oid,
                    picker_fio=assembly.get("picker_fio", ""),
                    assembly_start=start,
                    assembly_end=end,
                    assembly_duration_minutes=round(duration, 2) if duration is not None else None,
                    order_created_at=order.get("queue_time"),
                    wait_time_minutes=round(wait_minutes, 2) if wait_minutes is not None else None,
                    items_count=assembly.get("items_count"),
                    has_complaint=has_complaint,
                    complaint_count=complaint_count,
                    match_status="matched" if oid in orders_by_id else "unmatched",
                )
                db.session.add(ao)

                picker = assembly.get("picker_fio", "").strip()
                if not picker:
                    issues.append(
                        DataQualityIssue(
                            session_id=session.id,
                            issue_type="unknown_picker",
                            description=f'Не указан сборщик: ID "{oid}"',
                            source_sheet="assembly",
                            order_id=oid,
                            severity="warning",
                        )
                    )
                    picker = "(не указан)"

                picker_data[picker].append(
                    {
                        "order_id": oid,
                        "duration": duration,
                        "wait": wait_minutes,
                        "items": assembly.get("items_count"),
                        "has_complaint": has_complaint,
                        "complaint_count": complaint_count,
                        "start": start,
                        "end": end,
                    }
                )

        # Агрегация по сборщикам
        for picker_fio, records in picker_data.items():
            valid_durations = [
                r["duration"] for r in records if r["duration"] is not None and r["duration"] > 0
            ]
            valid_waits = [r["wait"] for r in records if r["wait"] is not None and r["wait"] >= 0]
            valid_items = [r["items"] for r in records if r["items"] is not None]
            complaints_records = [r for r in records if r["has_complaint"]]
            total_complaints = sum(r["complaint_count"] for r in records)
            above_threshold = sum(1 for r in valid_waits if r > wait_threshold)

            total_time = sum(valid_durations)
            total_items = sum(valid_items) if valid_items else 0

            pm = PickerMetrics(
                session_id=session.id,
                picker_fio=picker_fio,
                unique_orders=len(records),
                total_assembly_time_minutes=round(total_time, 2),
                avg_assembly_time_minutes=(
                    round(total_time / len(valid_durations), 2) if valid_durations else 0.0
                ),
                median_assembly_time_minutes=(
                    round(statistics.median(valid_durations), 2) if valid_durations else 0.0
                ),
                speed_orders_per_minute=(
                    round(len(valid_durations) / total_time, 4) if total_time > 0 else 0.0
                ),
                speed_items_per_minute=(
                    round(total_items / total_time, 4) if total_time > 0 and total_items else 0.0
                ),
                total_items=total_items,
                avg_wait_time_minutes=(
                    round(sum(valid_waits) / len(valid_waits), 2) if valid_waits else 0.0
                ),
                median_wait_time_minutes=(
                    round(statistics.median(valid_waits), 2) if valid_waits else 0.0
                ),
                min_wait_time_minutes=round(min(valid_waits), 2) if valid_waits else 0.0,
                max_wait_time_minutes=round(max(valid_waits), 2) if valid_waits else 0.0,
                orders_above_wait_threshold=above_threshold,
                orders_with_complaints=len(complaints_records),
                total_complaints=total_complaints,
                complaint_rate=(
                    round(len(complaints_records) / len(records) * 100, 1) if records else 0.0
                ),
            )
            db.session.add(pm)

        # Сохраняем проблемы качества данных
        for issue in issues:
            db.session.add(issue)

        session.error_count = sum(1 for i in issues if i.severity == "error")
        session.warning_count = sum(1 for i in issues if i.severity == "warning")
        session.status = "completed"

        db.session.commit()
        return session

    @classmethod
    def _detect_sheets_by_structure(cls, workbook, complaints_sheet, orders_sheet, assembly_sheet):
        """Определение листов по структуре столбцов, если названия не распознаны."""
        for ws in workbook.worksheets:
            if ws is complaints_sheet or ws is orders_sheet or ws is assembly_sheet:
                continue
            headers = [
                str(c.value).strip().lower() if c.value else ""
                for c in next(ws.iter_rows(min_row=1, max_row=1))
            ]
            header_text = " ".join(headers)
            if any(kw in header_text for kw in ["жалоб", "причин", "подтем"]):
                complaints_sheet = ws
            elif any(kw in header_text for kw in ["курьер", "адрес", "доставлен"]):
                orders_sheet = ws
            elif any(kw in header_text for kw in ["сбор", "собран"]):
                assembly_sheet = ws
        return complaints_sheet, orders_sheet, assembly_sheet

    @classmethod
    def _read_complaints_sheet(cls, sheet) -> tuple[list[dict], DataQualityIssue | None]:
        """Чтение листа жалоб."""
        if sheet is None:
            return [], DataQualityIssue(
                session_id=0,
                issue_type="sheet_not_found",
                description="Лист жалоб не найден",
                source_sheet="complaints",
                severity="error",
            )
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            return [], None

        headers = rows[0]
        cols = cls._detect_columns(
            headers,
            {
                "order_id": cls.ORDER_ID_KEYWORDS,
                "complaint_text": cls.COMPLAINT_TEXT_KEYWORDS,
            },
        )

        if "order_id" not in cols:
            return [], DataQualityIssue(
                session_id=0,
                issue_type="column_not_found",
                description="Не найден столбец ID заказа на листе жалоб",
                source_sheet="complaints",
                severity="error",
            )

        data = []
        for row in rows[1:]:
            if not any(row):
                continue
            oid = cls._normalize_id(row[cols["order_id"]])
            if not oid:
                continue
            complaint_text = ""
            if "complaint_text" in cols and cols["complaint_text"] < len(row):
                complaint_text = str(row[cols["complaint_text"]] or "").strip()
            data.append({"order_id": oid, "complaint_text": complaint_text})
        return data, None

    @classmethod
    def _read_orders_sheet(cls, sheet) -> tuple[list[dict], DataQualityIssue | None]:
        """Чтение листа заказов."""
        if sheet is None:
            return [], DataQualityIssue(
                session_id=0,
                issue_type="sheet_not_found",
                description="Лист заказов не найден",
                source_sheet="orders",
                severity="error",
            )
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            return [], None

        headers = rows[0]
        cols = cls._detect_columns(
            headers,
            {
                "order_id": cls.ORDER_ID_KEYWORDS,
                "queue_time": cls.QUEUE_TIME_KEYWORDS,
                "courier": cls.COURIER_KEYWORDS,
                "address": cls.ADDRESS_KEYWORDS,
            },
        )

        if "order_id" not in cols:
            return [], DataQualityIssue(
                session_id=0,
                issue_type="column_not_found",
                description="Не найден столбец ID заказа на листе заказов",
                source_sheet="orders",
                severity="error",
            )

        data = []
        for row in rows[1:]:
            if not any(row):
                continue
            oid = cls._normalize_id(row[cols["order_id"]])
            if not oid:
                continue
            queue_time = None
            if "queue_time" in cols and cols["queue_time"] < len(row):
                queue_time = cls._parse_datetime(row[cols["queue_time"]])
            courier = ""
            if "courier" in cols and cols["courier"] < len(row):
                courier = str(row[cols["courier"]] or "").strip()
            address = ""
            if "address" in cols and cols["address"] < len(row):
                address = str(row[cols["address"]] or "").strip()
            data.append(
                {
                    "order_id": oid,
                    "queue_time": queue_time,
                    "courier": courier,
                    "address": address,
                    "complaints": [],
                    "has_complaint": False,
                }
            )
        return data, None

    @classmethod
    def _read_assembly_sheet(cls, sheet) -> tuple[list[dict], DataQualityIssue | None]:
        """Чтение листа собранных заказов."""
        if sheet is None:
            return [], DataQualityIssue(
                session_id=0,
                issue_type="sheet_not_found",
                description="Лист собранных заказов не найден",
                source_sheet="assembly",
                severity="error",
            )
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            return [], None

        headers = rows[0]
        cols = cls._detect_columns(
            headers,
            {
                "order_id": cls.ORDER_ID_KEYWORDS,
                "picker_fio": cls.PICKER_KEYWORDS,
                "start_time": cls.START_TIME_KEYWORDS,
                "end_time": cls.END_TIME_KEYWORDS,
                "items_count": cls.ITEMS_KEYWORDS,
            },
        )

        if "order_id" not in cols:
            return [], DataQualityIssue(
                session_id=0,
                issue_type="column_not_found",
                description="Не найден столбец ID заказа на листе собранных заказов",
                source_sheet="assembly",
                severity="error",
            )

        data = []
        for row in rows[1:]:
            if not any(row):
                continue
            oid = cls._normalize_id(row[cols["order_id"]])
            if not oid:
                continue

            picker_fio = ""
            if "picker_fio" in cols and cols["picker_fio"] < len(row):
                picker_fio = str(row[cols["picker_fio"]] or "").strip()

            start = None
            if "start_time" in cols and cols["start_time"] < len(row):
                start = cls._parse_datetime(row[cols["start_time"]])

            end = None
            if "end_time" in cols and cols["end_time"] < len(row):
                end = cls._parse_datetime(row[cols["end_time"]])

            items_count = None
            if "items_count" in cols and cols["items_count"] < len(row):
                items_count = int(cls._parse_number(row[cols["items_count"]]) or 0) or None

            data.append(
                {
                    "order_id": oid,
                    "picker_fio": picker_fio,
                    "assembly_start": start,
                    "assembly_end": end,
                    "items_count": items_count,
                    "wait_time_minutes": None,
                }
            )
        return data, None

    @classmethod
    def get_session_results(cls, session_id: int) -> dict:
        """Получение результатов анализа для отображения."""
        session = db.session.get(OperationalAnalysisSession, session_id)
        if not session:
            return {}

        picker_metrics = (
            PickerMetrics.query.filter_by(session_id=session_id)
            .order_by(
                PickerMetrics.complaint_rate.desc(), PickerMetrics.speed_orders_per_minute.desc()
            )
            .all()
        )

        assembly_orders = (
            AssemblyOrder.query.filter_by(session_id=session_id)
            .order_by(AssemblyOrder.order_id)
            .all()
        )

        quality_issues = (
            DataQualityIssue.query.filter_by(session_id=session_id)
            .order_by(DataQualityIssue.severity, DataQualityIssue.issue_type)
            .all()
        )

        return {
            "session": session,
            "picker_metrics": picker_metrics,
            "assembly_orders": assembly_orders,
            "quality_issues": quality_issues,
            "filters": cls._get_filter_options(assembly_orders, picker_metrics),
        }

    @staticmethod
    def _get_filter_options(assembly_orders: list, picker_metrics: list) -> dict:
        """Получение опций фильтров."""
        pickers = sorted({m.picker_fio for m in picker_metrics})
        dates = sorted(
            {o.assembly_start.date().isoformat() for o in assembly_orders if o.assembly_start}
        )
        couriers = sorted(
            {o.courier_fio for o in assembly_orders if getattr(o, "courier_fio", None)}
        )
        return {
            "pickers": pickers,
            "dates": dates,
            "couriers": couriers,
            "has_complaint": any(o.has_complaint for o in assembly_orders),
        }

    @classmethod
    def filter_results(
        cls,
        session_id: int,
        picker: str = "",
        date_from: str = "",
        date_to: str = "",
        time_of_day: str = "",
        has_complaint: str = "",
    ) -> dict:
        """Фильтрация результатов по сборщику, дате, времени суток, жалобе."""
        results = cls.get_session_results(session_id)
        if not results:
            return {}

        orders = results["assembly_orders"]
        pickers = results["picker_metrics"]

        if picker:
            orders = [o for o in orders if o.picker_fio == picker]
            pickers = [p for p in pickers if p.picker_fio == picker]

        if date_from:
            df = datetime.strptime(date_from, "%Y-%m-%d").date()
            orders = [o for o in orders if o.assembly_start and o.assembly_start.date() >= df]

        if date_to:
            dt = datetime.strptime(date_to, "%Y-%m-%d").date()
            orders = [o for o in orders if o.assembly_start and o.assembly_start.date() <= dt]

        if time_of_day == "morning":
            orders = [o for o in orders if o.assembly_start and 6 <= o.assembly_start.hour < 12]
        elif time_of_day == "afternoon":
            orders = [o for o in orders if o.assembly_start and 12 <= o.assembly_start.hour < 18]
        elif time_of_day == "evening":
            orders = [o for o in orders if o.assembly_start and 18 <= o.assembly_start.hour < 24]
        elif time_of_day == "night":
            orders = [o for o in orders if o.assembly_start and 0 <= o.assembly_start.hour < 6]

        if has_complaint == "yes":
            orders = [o for o in orders if o.has_complaint]
        elif has_complaint == "no":
            orders = [o for o in orders if not o.has_complaint]

        results["assembly_orders"] = orders
        results["picker_metrics"] = pickers
        return results

    @classmethod
    def export_excel(cls, session_id: int) -> bytes:
        """Экспорт результатов анализа в Excel."""
        results = cls.get_session_results(session_id)
        if not results:
            return b""

        results["session"]
        wb = Workbook()

        # Сводка по сборщикам
        ws1 = wb.active
        ws1.title = "Сводка по сборщикам"
        ws1.append(
            [
                "Сборщик",
                "Уникальных заказов",
                "Среднее время сборки (мин)",
                "Медианное время сборки (мин)",
                "Скорость (заказов/мин)",
                "Скорость (товаров/мин)",
                "Среднее ожидание (мин)",
                "Медианное ожидание (мин)",
                "Заказов с превышением ожидания",
                "Заказов с жалобами",
                "Всего жалоб",
                "Доля заказов с жалобами (%)",
            ]
        )
        for pm in results["picker_metrics"]:
            ws1.append(
                [
                    pm.picker_fio,
                    pm.unique_orders,
                    pm.avg_assembly_time_minutes,
                    pm.median_assembly_time_minutes,
                    pm.speed_orders_per_minute,
                    pm.speed_items_per_minute,
                    pm.avg_wait_time_minutes,
                    pm.median_wait_time_minutes,
                    pm.orders_above_wait_threshold,
                    pm.orders_with_complaints,
                    pm.total_complaints,
                    pm.complaint_rate,
                ]
            )

        # Детализация заказов
        ws2 = wb.create_sheet("Детализация заказов")
        ws2.append(
            [
                "ID заказа",
                "Дата и время",
                "Сборщик",
                "Время начала сборки",
                "Время окончания сборки",
                "Продолжительность сборки (мин)",
                "Ожидание сборки (мин)",
                "Количество товаров",
                "Наличие жалобы",
                "Количество жалоб",
                "Статус сопоставления",
            ]
        )
        for ao in results["assembly_orders"]:
            ws2.append(
                [
                    ao.order_id,
                    ao.assembly_start.strftime("%Y-%m-%d %H:%M:%S") if ao.assembly_start else "",
                    ao.picker_fio,
                    ao.assembly_start.strftime("%H:%M:%S") if ao.assembly_start else "",
                    ao.assembly_end.strftime("%H:%M:%S") if ao.assembly_end else "",
                    (
                        ao.assembly_duration_minutes
                        if ao.assembly_duration_minutes is not None
                        else ""
                    ),
                    ao.wait_time_minutes if ao.wait_time_minutes is not None else "",
                    ao.items_count or "",
                    "Да" if ao.has_complaint else "Нет",
                    ao.complaint_count,
                    ao.match_status,
                ]
            )

        # Контроль качества
        ws3 = wb.create_sheet("Контроль качества")
        ws3.append(
            [
                "Тип проблемы",
                "Описание",
                "Источник",
                "Строка",
                "ID заказа",
                "Сборщик",
                "Критичность",
            ]
        )
        for qi in results["quality_issues"]:
            ws3.append(
                [
                    qi.issue_type,
                    qi.description,
                    qi.source_sheet,
                    qi.row_number or "",
                    qi.order_id or "",
                    qi.picker_fio or "",
                    qi.severity,
                ]
            )

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @classmethod
    def get_recent_sessions(cls, limit: int = 20) -> list:
        """Получение последних сессий анализа."""
        return (
            OperationalAnalysisSession.query.order_by(OperationalAnalysisSession.uploaded_at.desc())
            .limit(limit)
            .all()
        )

    @classmethod
    def delete_session(cls, session_id: int) -> bool:
        """Удаление сессии анализа со всеми связанными данными."""
        session = db.session.get(OperationalAnalysisSession, session_id)
        if not session:
            return False
        db.session.delete(session)
        db.session.commit()
        return True
