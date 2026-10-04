import os
import logging
from logging.handlers import RotatingFileHandler
from datetime import datetime
from pathlib import Path

from flask import (
    Flask, render_template, request, redirect, url_for,
    flash, session, send_from_directory, abort, jsonify
)
from werkzeug.utils import secure_filename
from werkzeug.security import generate_password_hash, check_password_hash
import openpyxl

from config import (
    SECRET_KEY, ADMIN_PASSWORD, ADMIN_TOKEN, DATABASE_PATH,
    UPLOAD_FOLDER, MAX_CONTENT_LENGTH, ALLOWED_EXTENSIONS,
    RATELIMIT_STORAGE_URL, FLASK_DEBUG,
    MAIL_SERVER, MAIL_PORT, MAIL_USE_TLS, MAIL_USE_SSL,
    MAIL_USERNAME, MAIL_PASSWORD, MAIL_DEFAULT_SENDER, MAIL_ADMIN_RECIPIENTS
)
from models import db, User, Order, Photo, generate_token, ImportMapping
from flask_wtf.csrf import CSRFProtect
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
from flask_migrate import Migrate
from flask_mail import Mail, Message

app = Flask(__name__)
app.config['SECRET_KEY'] = SECRET_KEY
app.config['MAX_CONTENT_LENGTH'] = MAX_CONTENT_LENGTH
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['SQLALCHEMY_DATABASE_URI'] = f'sqlite:///{DATABASE_PATH}'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

# Mail config
app.config['MAIL_SERVER'] = MAIL_SERVER
app.config['MAIL_PORT'] = MAIL_PORT
app.config['MAIL_USE_TLS'] = MAIL_USE_TLS
app.config['MAIL_USE_SSL'] = MAIL_USE_SSL
app.config['MAIL_USERNAME'] = MAIL_USERNAME
app.config['MAIL_PASSWORD'] = MAIL_PASSWORD
app.config['MAIL_DEFAULT_SENDER'] = MAIL_DEFAULT_SENDER

Path(UPLOAD_FOLDER).mkdir(parents=True, exist_ok=True)

db.init_app(app)
migrate = Migrate(app, db)
csrf = CSRFProtect(app)
mail = Mail(app)

limiter = Limiter(
    get_remote_address,
    app=app,
    storage_uri=RATELIMIT_STORAGE_URL,
    default_limits=["200 per day", "50 per hour"]
)

# Email functions
def send_email(subject, recipients, html_body, text_body=None, attachments=None):
    """Send email with optional attachments."""
    if not MAIL_USERNAME or not MAIL_PASSWORD:
        app.logger.warning('Email not configured, skipping send')
        return False
    
    try:
        msg = Message(
            subject=subject,
            recipients=recipients,
            sender=MAIL_DEFAULT_SENDER,
            html=html_body,
            body=text_body or ''
        )
        if attachments:
            for filename, content, mimetype in attachments:
                msg.attach(filename, mimetype, content)
        mail.send(msg)
        app.logger.info(f'Email sent to {recipients}: {subject}')
        return True
    except Exception as e:
        app.logger.error(f'Failed to send email: {e}')
        return False


def send_order_completion_email(order, user):
    """Send email when order is completed with photos."""
    if not MAIL_ADMIN_RECIPIENTS:
        return False
    
    photos = Photo.query.filter_by(order_id=order.id).all()
    
    html = render_template('emails/order_completed.html', 
                           order=order, user=user, photos=photos)
    text = f'''
Заказ подтвержден: {order.order_number}
Клиент: {user.fio}
Описание: {order.description}
Фото: {len(photos)} шт.
Дата: {order.created_at}
'''
    
    attachments = []
    for photo in photos:
        filepath = os.path.join(UPLOAD_FOLDER, photo.filename)
        if os.path.exists(filepath):
            with open(filepath, 'rb') as f:
                attachments.append((
                    photo.original_filename or photo.filename,
                    f.read(),
                    'image/jpeg'
                ))
    
    return send_email(
        subject=f'✅ Заказ #{order.order_number} подтвержден — {len(photos)} фото',
        recipients=MAIL_ADMIN_RECIPIENTS,
        html_body=html,
        text_body=text,
        attachments=attachments if attachments else None
    )


def send_order_created_email(order, user):
    """Send notification when new order is created."""
    if not MAIL_ADMIN_RECIPIENTS:
        return False
    
    html = render_template('emails/order_created.html', order=order, user=user)
    text = f'''
Новый заказ: {order.order_number}
Клиент: {user.fio}
Описание: {order.description}
Дата: {order.created_at}
'''
    
    return send_email(
        subject=f'📦 Новый заказ #{order.order_number}',
        recipients=MAIL_ADMIN_RECIPIENTS,
        html_body=html,
        text_body=text
    )

if not app.debug:
    log_dir = Path('logs')
    log_dir.mkdir(exist_ok=True)
    file_handler = RotatingFileHandler(
        log_dir / 'app.log',
        maxBytes=10240,
        backupCount=10
    )
    file_handler.setFormatter(logging.Formatter(
        '%(asctime)s %(levelname)s: %(message)s [in %(pathname)s:%(lineno)d]'
    ))
    file_handler.setLevel(logging.INFO)
    app.logger.addHandler(file_handler)
    app.logger.setLevel(logging.INFO)
    app.logger.info('GL_JB startup')

# Create database tables on startup
with app.app_context():
    try:
        db.create_all()
        app.logger.info('Database tables created/verified')
        # Verify tables exist
        from sqlalchemy import inspect
        inspector = inspect(db.engine)
        tables = inspector.get_table_names()
        app.logger.info(f'Tables in database: {tables}')
    except Exception as e:
        app.logger.error(f'Database initialization error: {e}')

def allowed_file(filename):
    return '.' in filename and \
           filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def admin_required(f):
    from functools import wraps
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('is_admin'):
            return redirect(url_for('admin_login'))
        return f(*args, **kwargs)
    return decorated_function

def user_required(f):
    from functools import wraps
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('user_id'):
            return redirect(url_for('unified_login'))
        return f(*args, **kwargs)
    return decorated_function

@app.route('/admin/login', methods=['GET', 'POST'])
@limiter.limit("5 per minute")
def admin_login():
    if request.method == 'POST':
        password = request.form.get('password', '')
        if password == ADMIN_PASSWORD:
            session['is_admin'] = True
            app.logger.info(f'Admin login from {get_remote_address()}')
            return redirect(url_for('admin_dashboard'))
        flash('Неверный пароль', 'error')
        app.logger.warning(f'Failed admin login from {get_remote_address()}')
    return render_template('admin_login.html')

@app.route('/admin/logout')
def admin_logout():
    session.clear()
    return redirect(url_for('admin_login'))

@app.route('/admin/token_login', methods=['GET', 'POST'])
@limiter.limit("5 per minute")
def admin_token_login():
    if request.method == 'POST':
        token = request.form.get('token', '').strip()
        if token == ADMIN_TOKEN:
            session['is_admin'] = True
            app.logger.info(f'Admin token login from {get_remote_address()}')
            return redirect(url_for('admin_dashboard'))
        flash('Неверный токен суперадмина', 'error')
        app.logger.warning(f'Failed admin token login from {get_remote_address()}')
    return render_template('admin_token_login.html')

@app.route('/admin')
@admin_required
def admin_dashboard():
    users_count = User.query.count()
    orders_count = Order.query.count()
    photos_count = Photo.query.count()
    
    # Orders with photos (confirmed)
    orders_with_photos = db.session.query(Order.id).join(Photo, Order.id == Photo.order_id).distinct().count()
    
    # Upload folder info
    upload_folder = os.path.abspath(UPLOAD_FOLDER)
    upload_size = 0
    if os.path.exists(upload_folder):
        for root, dirs, files in os.walk(upload_folder):
            for f in files:
                fp = os.path.join(root, f)
                if os.path.isfile(fp):
                    upload_size += os.path.getsize(fp)
    
    return render_template('admin_dashboard.html',
                           users_count=users_count,
                           orders_count=orders_count,
                           photos_count=photos_count,
                           orders_with_photos=orders_with_photos,
                           upload_folder=upload_folder,
                           upload_size=upload_size)

@app.route('/admin/users', methods=['GET', 'POST'])
@admin_required
def admin_users():
    page = request.args.get('page', 1, type=int)
    search = request.args.get('search', '', type=str)
    
    query = User.query
    if search:
        query = query.filter(User.fio.ilike(f'%{search}%'))
    
    pagination = query.order_by(User.created_at.desc()).paginate(
        page=page, per_page=20, error_out=False
    )
    users = pagination.items
    
    if request.method == 'POST':
        fio = request.form.get('fio', '').strip()
        if fio:
            token = generate_token()
            try:
                user = User(fio=fio, token=token)
                db.session.add(user)
                db.session.commit()
                flash(f'Пользователь создан. Токен: {token}', 'success')
                app.logger.info(f'Created user: {fio} with token {token}')
            except Exception as e:
                db.session.rollback()
                flash('Пользователь с таким ФИО уже существует', 'error')
                app.logger.error(f'Failed to create user {fio}: {e}')
        else:
            flash('Введите ФИО', 'error')
        return redirect(url_for('admin_users'))
    
    return render_template('admin_users.html', users=users, pagination=pagination, search=search)

@app.route('/admin/users/delete/<int:user_id>', methods=['POST'])
@admin_required
def admin_delete_user(user_id):
    user = User.query.get_or_404(user_id)
    fio = user.fio
    db.session.delete(user)
    db.session.commit()
    flash('Пользователь удалён', 'success')
    app.logger.info(f'Deleted user: {fio}')
    return redirect(url_for('admin_users'))

@app.route('/admin/users/regenerate_token/<int:user_id>', methods=['POST'])
@admin_required
def admin_regenerate_token(user_id):
    user = User.query.get_or_404(user_id)
    old_token = user.token
    user.token = generate_token()
    db.session.commit()
    flash(f'Токен обновлён. Новый токен: {user.token}', 'success')
    app.logger.info(f'Regenerated token for user {user.fio}: {old_token} -> {user.token}')
    return redirect(url_for('admin_users'))

@app.route('/admin/upload_excel', methods=['GET', 'POST'])
@admin_required
def admin_upload_excel():
    if request.method == 'POST':
        file = request.files.get('excel_file')
        if not file or file.filename == '':
            flash('Файл не выбран', 'error')
            return redirect(url_for('admin_upload_excel'))
        
        if not file.filename.endswith(('.xlsx', '.xls')):
            flash('Неверный формат файла. Нужен .xlsx или .xls', 'error')
            return redirect(url_for('admin_upload_excel'))
        
        try:
            workbook = openpyxl.load_workbook(file)
            sheet = workbook.active
            
            headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            
            fio_col_idx = None
            fio_variants = ['фио', 'fio', 'ФИО', 'Фио', 'Фамилия Имя Отчество', 'ФИО клиента']
            for idx, header in enumerate(headers):
                if header and str(header).strip().lower() in [v.lower() for v in fio_variants]:
                    fio_col_idx = idx
                    break
            
            if fio_col_idx is None:
                flash('Не найдена колонка с ФИО. Ожидаемые названия: "ФИО", "fio", "ФИО клиента"', 'error')
                return redirect(url_for('admin_upload_excel'))
            
            users_map = {u.fio.strip().lower(): u.id for u in User.query.all()}
            
            imported = 0
            skipped = 0
            errors = []
            duplicates = 0
            
            for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if not any(row):
                    continue
                
                fio_value = row[fio_col_idx]
                if not fio_value:
                    skipped += 1
                    errors.append(f'Строка {row_idx}: пустое ФИО')
                    continue
                
                fio_normalized = str(fio_value).strip().lower()
                user_id = users_map.get(fio_normalized)
                
                if not user_id:
                    skipped += 1
                    errors.append(f'Строка {row_idx}: пользователь "{fio_value}" не найден')
                    continue
                
                order_data = {}
                for idx, header in enumerate(headers):
                    if idx != fio_col_idx and header:
                        order_data[str(header).strip()] = row[idx]
                
                order_number = order_data.get('Номер заказа') or order_data.get('order_number') or f'Заказ #{row_idx}'
                
                existing = Order.query.filter_by(user_id=user_id, order_number=order_number).first()
                if existing:
                    duplicates += 1
                    errors.append(f'Строка {row_idx}: заказ "{order_number}" уже существует для этого пользователя')
                    continue
                
                description_parts = []
                for key, value in order_data.items():
                    if value is not None:
                        description_parts.append(f'{key}: {value}')
                description = '; '.join(description_parts)
                
                order = Order(user_id=user_id, order_number=order_number, description=description)
                db.session.add(order)
                imported += 1
            
            db.session.commit()
            
            msg = f'Импортировано заказов: {imported}. Пропущено строк: {skipped}. Дублей: {duplicates}.'
            if errors:
                msg += ' Ошибки: ' + '; '.join(errors[:5])
                if len(errors) > 5:
                    msg += f' ... и ещё {len(errors) - 5} ошибок.'
            flash(msg, 'success' if imported > 0 else 'warning')
            app.logger.info(f'Excel import: {imported} imported, {skipped} skipped, {duplicates} duplicates')
            
        except Exception as e:
            db.session.rollback()
            flash(f'Ошибка при обработке файла: {str(e)}', 'error')
            app.logger.error(f'Excel import error: {e}')
        
        return redirect(url_for('admin_upload_excel'))
    
    return render_template('admin_upload_excel.html')

# Import Mappings
@app.route('/admin/import_mappings', methods=['GET', 'POST'])
@admin_required
def admin_import_mappings():
    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        type_ = request.form.get('type', '').strip()
        id_column = request.form.get('id_column', '').strip()
        fio_column = request.form.get('fio_column', '').strip()
        order_number_column = request.form.get('order_number_column', '').strip()
        
        # Parse field mapping JSON
        field_mapping = {}
        keys = request.form.getlist('mapping_key[]')
        values = request.form.getlist('mapping_value[]')
        for k, v in zip(keys, values):
            if k.strip() and v.strip():
                field_mapping[k.strip()] = v.strip()
        
        if not name or not type_ or not id_column:
            flash('Заполните обязательные поля: название, тип, колонка ID', 'error')
        else:
            try:
                mapping = ImportMapping(
                    name=name,
                    type=type_,
                    id_column=id_column,
                    fio_column=fio_column or None,
                    order_number_column=order_number_column or None,
                    field_mapping=field_mapping
                )
                db.session.add(mapping)
                db.session.commit()
                flash('Маппинг создан', 'success')
            except Exception as e:
                db.session.rollback()
                flash(f'Ошибка: {e}', 'error')
        return redirect(url_for('admin_import_mappings'))
    
    mappings = ImportMapping.query.order_by(ImportMapping.created_at.desc()).all()
    return render_template('admin_import_mappings.html', mappings=mappings)


@app.route('/admin/import_mappings/delete/<int:mapping_id>', methods=['POST'])
@admin_required
def admin_delete_import_mapping(mapping_id):
    mapping = ImportMapping.query.get_or_404(mapping_id)
    db.session.delete(mapping)
    db.session.commit()
    flash('Маппинг удалён', 'success')
    return redirect(url_for('admin_import_mappings'))


@app.route('/admin/import_excel', methods=['GET', 'POST'])
@admin_required
def admin_import_excel():
    mappings = ImportMapping.query.order_by(ImportMapping.type, ImportMapping.name).all()
    
    if request.method == 'POST':
        mapping_id = request.form.get('mapping_id', type=int)
        file = request.files.get('excel_file')
        
        if not mapping_id:
            flash('Выберите маппинг', 'error')
            return redirect(url_for('admin_import_excel'))
        
        if not file or file.filename == '':
            flash('Файл не выбран', 'error')
            return redirect(url_for('admin_import_excel'))
        
        if not file.filename.endswith(('.xlsx', '.xls')):
            flash('Неверный формат файла. Нужен .xlsx или .xls', 'error')
            return redirect(url_for('admin_import_excel'))
        
        mapping = ImportMapping.query.get_or_404(mapping_id)
        
        try:
            workbook = openpyxl.load_workbook(file)
            sheet = workbook.active
            
            headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            headers = [str(h).strip() if h else '' for h in headers]
            
            # Find column indices
            id_col_idx = None
            for idx, header in enumerate(headers):
                if header.lower() == mapping.id_column.lower():
                    id_col_idx = idx
                    break
            
            if id_col_idx is None:
                flash(f'Колонка ID "{mapping.id_column}" не найдена в файле. Доступные: {headers}', 'error')
                return redirect(url_for('admin_import_excel'))
            
            fio_col_idx = None
            if mapping.fio_column:
                for idx, header in enumerate(headers):
                    if header.lower() == mapping.fio_column.lower():
                        fio_col_idx = idx
                        break
            
            order_num_col_idx = None
            if mapping.order_number_column:
                for idx, header in enumerate(headers):
                    if header.lower() == mapping.order_number_column.lower():
                        order_num_col_idx = idx
                        break
            
            # Build field mapping indices
            field_indices = {}
            for file_col, model_field in mapping.field_mapping.items():
                for idx, header in enumerate(headers):
                    if header.lower() == file_col.lower():
                        field_indices[model_field] = idx
                        break
            
            users_map = {u.fio.strip().lower(): u.id for u in User.query.all()}
            
            imported = 0
            updated = 0
            skipped = 0
            errors = []
            
            start_row = 2 if mapping.skip_first_row else 1
            
            for row_idx, row in enumerate(sheet.iter_rows(min_row=start_row, values_only=True), start=start_row):
                if not any(row):
                    continue
                
                ext_id = row[id_col_idx]
                if not ext_id:
                    skipped += 1
                    errors.append(f'Строка {row_idx}: пустой ID')
                    continue
                
                ext_id = str(ext_id).strip()
                
                if mapping.type == 'orders':
                    # Import orders
                    fio_value = row[fio_col_idx] if fio_col_idx is not None else None
                    if not fio_value:
                        skipped += 1
                        errors.append(f'Строка {row_idx}: пустое ФИО')
                        continue
                    
                    fio_normalized = str(fio_value).strip().lower()
                    user_id = users_map.get(fio_normalized)
                    if not user_id:
                        skipped += 1
                        errors.append(f'Строка {row_idx}: пользователь "{fio_value}" не найден')
                        continue
                    
                    order_number = row[order_num_col_idx] if order_num_col_idx is not None else f'Заказ #{row_idx}'
                    if not order_number:
                        order_number = f'Заказ #{row_idx}'
                    
                    # Check if order with this external_id exists
                    order = Order.query.filter_by(external_id=ext_id).first()
                    
                    description_parts = []
                    for model_field, col_idx in field_indices.items():
                        value = row[col_idx]
                        if value is not None:
                            description_parts.append(f'{model_field}: {value}')
                    description = '; '.join(description_parts)
                    
                    if order:
                        order.user_id = user_id
                        order.order_number = order_number
                        order.description = description
                        order.external_id = ext_id
                        updated += 1
                    else:
                        order = Order(
                            user_id=user_id,
                            order_number=order_number,
                            description=description,
                            external_id=ext_id
                        )
                        db.session.add(order)
                        imported += 1
                
                elif mapping.type == 'complaints':
                    # Import complaints - link to existing order by external_id
                    order = Order.query.filter_by(external_id=ext_id).first()
                    if not order:
                        skipped += 1
                        errors.append(f'Строка {row_idx}: заказ с ID "{ext_id}" не найден')
                        continue
                    
                    # Update complaint fields
                    for model_field, col_idx in field_indices.items():
                        value = row[col_idx]
                        if value is not None:
                            if model_field == 'complaint_text':
                                order.complaint_text = str(value)
                            elif model_field == 'complaint_status':
                                order.complaint_status = str(value)
                            elif model_field == 'complaint_date':
                                if isinstance(value, str):
                                    try:
                                        from datetime import datetime
                                        order.complaint_date = datetime.fromisoformat(value.replace('Z', '+00:00'))
                                    except:
                                        pass
                                elif hasattr(value, 'isoformat'):
                                    order.complaint_date = value
                    
                    updated += 1
            
            db.session.commit()
            
            msg = f'Импорт завершён. Создано: {imported}, обновлено: {updated}, пропущено: {skipped}.'
            if errors:
                msg += ' Ошибки: ' + '; '.join(errors[:5])
                if len(errors) > 5:
                    msg += f' ... и ещё {len(errors) - 5} ошибок.'
            flash(msg, 'success' if (imported + updated) > 0 else 'warning')
            app.logger.info(f'Excel import ({mapping.type}): {imported} created, {updated} updated, {skipped} skipped')
            
        except Exception as e:
            db.session.rollback()
            flash(f'Ошибка при обработке файла: {str(e)}', 'error')
            app.logger.error(f'Excel import error: {e}')
        
        return redirect(url_for('admin_import_excel'))
    
    return render_template('admin_import_excel.html', mappings=mappings)

@app.route('/admin/orders')
@admin_required
def admin_orders():
    page = request.args.get('page', 1, type=int)
    search = request.args.get('search', '', type=str)
    user_filter = request.args.get('user_id', '', type=str)
    photo_filter = request.args.get('photo_filter', '', type=str)
    
    query = db.session.query(Order, User).join(User, Order.user_id == User.id)
    
    if search:
        query = query.filter(
            db.or_(
                Order.order_number.ilike(f'%{search}%'),
                Order.description.ilike(f'%{search}%'),
                User.fio.ilike(f'%{search}%')
            )
        )
    
    if user_filter:
        query = query.filter(User.id == user_filter)
    
    if photo_filter == 'with_photos':
        query = query.filter(Order.photos.any())
    elif photo_filter == 'without_photos':
        query = query.filter(~Order.photos.any())
    
    pagination = query.order_by(Order.created_at.desc()).paginate(
        page=page, per_page=20, error_out=False
    )
    
    orders_with_users = []
    for order, user in pagination.items:
        orders_with_users.append({
            'order': order,
            'user': user,
            'photos_count': len(order.photos)
        })
    
    users = User.query.order_by(User.fio).all()
    
    return render_template('admin_orders.html', 
                           orders=orders_with_users, 
                           pagination=pagination, 
                           search=search,
                           users=users,
                           user_filter=user_filter,
                           photo_filter=photo_filter)


@app.route('/admin/orders/send_email/<int:order_id>', methods=['POST'])
@admin_required
def admin_send_order_email(order_id):
    order = Order.query.get_or_404(order_id)
    user = User.query.get(order.user_id)
    
    if not user:
        flash('Пользователь не найден', 'error')
        return redirect(url_for('admin_orders'))
    
    if send_order_completion_email(order, user):
        flash('Email отправлен администраторам', 'success')
    else:
        flash('Ошибка отправки (проверьте настройки почты)', 'error')
    
    return redirect(url_for('admin_orders'))

@app.route('/admin/orders/delete/<int:order_id>', methods=['POST'])
@admin_required
def admin_delete_order(order_id):
    order = Order.query.get_or_404(order_id)
    db.session.delete(order)
    db.session.commit()
    flash('Заказ удалён', 'success')
    app.logger.info(f'Deleted order: {order_id}')
    return redirect(url_for('admin_orders'))

@app.route('/admin/orders/edit/<int:order_id>', methods=['GET', 'POST'])
@admin_required
def admin_edit_order(order_id):
    order = Order.query.get_or_404(order_id)
    
    if request.method == 'POST':
        order.order_number = request.form.get('order_number', '').strip()
        order.description = request.form.get('description', '').strip()
        db.session.commit()
        flash('Заказ обновлён', 'success')
        app.logger.info(f'Edited order: {order_id}')
        return redirect(url_for('admin_orders'))
    
    return render_template('admin_edit_order.html', order=order)

@app.route('/admin/export_excel')
@admin_required
def admin_export_excel():
    from openpyxl import Workbook
    from flask import send_file
    import io
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Заказы"
    
    headers = ['ID', 'ФИО', 'Токен', 'Номер заказа', 'Описание', 'Создан', 'Кол-во фото']
    ws.append(headers)
    
    orders = db.session.query(Order, User).join(User).order_by(Order.created_at.desc()).all()
    
    for order, user in orders:
        ws.append([
            order.id,
            user.fio,
            user.token,
            order.order_number,
            order.description,
            order.created_at.strftime('%Y-%m-%d %H:%M:%S') if order.created_at else '',
            len(order.photos)
        ])
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'orders_export_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    )

@app.route('/login', methods=['GET', 'POST'])
@limiter.limit("10 per minute")
def unified_login():
    if request.method == 'POST':
        token = request.form.get('token', '').strip()
        
        if not token:
            flash('Введите токен', 'error')
            return render_template('unified_login.html')
        
        # Суперадмин по токену
        if token == ADMIN_TOKEN:
            session['is_admin'] = True
            app.logger.info(f'Superadmin token login from {get_remote_address()}')
            return redirect(url_for('admin_dashboard'))
        
        # Обычный пользователь по токену
        user = User.query.filter_by(token=token).first()
        if user:
            session['user_id'] = user.id
            session['user_fio'] = user.fio
            app.logger.info(f'User login: {user.fio} from {get_remote_address()}')
            return redirect(url_for('user_orders'))
        else:
            flash('Неверный токен', 'error')
            app.logger.warning(f'Failed login with token {token} from {get_remote_address()}')
    
    return render_template('unified_login.html')

@app.route('/logout')
def user_logout():
    session.clear()
    return redirect(url_for('unified_login'))

@app.route('/orders')
@user_required
def user_orders():
    user_id = session['user_id']
    page = request.args.get('page', 1, type=int)
    search = request.args.get('search', '', type=str)
    
    query = Order.query.filter_by(user_id=user_id)
    
    if search:
        query = query.filter(
            db.or_(
                Order.order_number.ilike(f'%{search}%'),
                Order.description.ilike(f'%{search}%')
            )
        )
    
    pagination = query.order_by(Order.created_at.desc()).paginate(
        page=page, per_page=10, error_out=False
    )
    orders = pagination.items
    
    for order in orders:
        order.photos_count = len(order.photos)
    
    return render_template('user_orders.html', orders=orders, pagination=pagination, search=search)

@app.route('/orders/<int:order_id>/upload', methods=['POST'])
@user_required
def upload_photo(order_id):
    user_id = session['user_id']
    
    order = Order.query.filter_by(id=order_id, user_id=user_id).first()
    
    if not order:
        flash('Заказ не найден', 'error')
        return redirect(url_for('user_orders'))
    
    if 'photo' not in request.files:
        flash('Файл не выбран', 'error')
        return redirect(url_for('user_orders'))
    
    file = request.files['photo']
    if file.filename == '':
        flash('Файл не выбран', 'error')
        return redirect(url_for('user_orders'))
    
    if not allowed_file(file.filename):
        flash('Недопустимый формат. Разрешены: jpg, png, webp', 'error')
        return redirect(url_for('user_orders'))
    
    original_filename = secure_filename(file.filename)
    ext = original_filename.rsplit('.', 1)[1].lower()
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f'{order_id}_{timestamp}.{ext}'
    filepath = os.path.join(UPLOAD_FOLDER, filename)
    file.save(filepath)
    
    photo = Photo(order_id=order_id, filename=filename, original_filename=original_filename)
    db.session.add(photo)
    db.session.commit()
    
    # Send email notification to admins
    user = User.query.get(user_id)
    if user:
        send_order_completion_email(order, user)
    
    flash('Фото загружено', 'success')
    app.logger.info(f'Photo uploaded: {filename} for order {order_id} by user {user_id}')
    return redirect(url_for('user_orders'))

@app.route('/orders/<int:order_id>/photos')
@user_required
def order_photos(order_id):
    user_id = session['user_id']
    
    order = Order.query.filter_by(id=order_id, user_id=user_id).first()
    
    if not order:
        flash('Заказ не найден', 'error')
        return redirect(url_for('user_orders'))
    
    photos = Photo.query.filter_by(order_id=order_id).order_by(Photo.uploaded_at.desc()).all()
    
    return render_template('user_order_photos.html', order=order, photos=photos)

@app.route('/orders/<int:order_id>/photos/delete/<int:photo_id>', methods=['POST'])
@user_required
def delete_photo(order_id, photo_id):
    user_id = session['user_id']
    
    order = Order.query.filter_by(id=order_id, user_id=user_id).first()
    if not order:
        flash('Заказ не найден', 'error')
        return redirect(url_for('user_orders'))
    
    photo = Photo.query.filter_by(id=photo_id, order_id=order_id).first()
    if not photo:
        flash('Фото не найдено', 'error')
        return redirect(url_for('order_photos', order_id=order_id))
    
    filepath = os.path.join(UPLOAD_FOLDER, photo.filename)
    if os.path.exists(filepath):
        os.remove(filepath)
    
    db.session.delete(photo)
    db.session.commit()
    
    flash('Фото удалено', 'success')
    app.logger.info(f'Photo deleted: {photo.filename} from order {order_id} by user {user_id}')
    return redirect(url_for('order_photos', order_id=order_id))

@app.route('/uploads/<filename>')
def uploaded_file(filename):
    return send_from_directory(UPLOAD_FOLDER, filename)

@app.route('/')
def index():
    return redirect(url_for('unified_login'))

@app.route('/health')
def health_check():
    return jsonify({'status': 'ok', 'timestamp': datetime.utcnow().isoformat()})

if __name__ == '__main__':
    with app.app_context():
        db.create_all()
    app.run(host='0.0.0.0', port=5000, debug=FLASK_DEBUG)