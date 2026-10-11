"""Regression: статические правки P1, которые легко откатить «по ошибке».

- мобильные таблицы больше не теряют колонку «Действия»;
- nginx не кэширует приватные /uploads/ как public, immutable;
- service worker не кладёт приватный HTML в Cache Storage;
- Excel-выгрузка не отдаёт токен пользователя (это пароль для входа);
- healthcheck в docker-compose не использует curl (его нет в python:3.11-slim).
"""

import io
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]


def _read(relative: str) -> str:
    return (PROJECT_ROOT / relative).read_text(encoding="utf-8")


def test_mobile_tables_do_not_hide_columns():
    css = _read("app/static/style.css")

    assert "nth-child(n+4)" not in css, "колонки таблиц снова скрываются на мобильных"
    assert ".table-wrap" in css, "горизонтальный скролл таблиц удалён"
    assert "overflow-x: auto" in css


def test_mobile_tables_can_switch_to_card_view():
    """P3-3: вместо скролла колонок — карточки с подписями из thead."""
    css = _read("app/static/style.css")
    html = _read("app/templates/base.html")

    assert ".is-cards" in css, "нет стилей карточного вида таблиц"
    assert "attr(data-label)" in css, "значения ячеек не подписаны заголовками"
    assert "@media (max-width: 768px)" in css

    assert "is-cards" in html, "JS не включает карточный вид на мобильных"
    assert "data-label" in html, "заголовки колонок не копируются в data-label"
    # Деградация без JS: карточки включаются только классом,
    # обычная таблица со скроллом должна остаться рабочей
    assert ".table-wrap table.is-cards" in css


def test_nginx_does_not_cache_private_uploads():
    conf = _read("nginx.conf")

    assert "public, immutable" not in conf, "/uploads/ снова кэшируется как публичные"
    assert "private, no-store" in conf


def test_app_healthcheck_does_not_use_curl():
    import yaml

    compose = yaml.safe_load(_read("docker-compose.yml"))
    healthcheck = compose["services"]["app"]["healthcheck"]["test"]
    joined = " ".join(healthcheck)

    assert "curl" not in joined, "python:3.11-slim не содержит curl — healthcheck всегда unhealthy"
    assert healthcheck[0] == "CMD"
    assert "urllib.request" in joined


def test_upload_form_has_no_xss_and_no_alert():
    """P2-1: имя файла попадало в innerHTML → XSS; P2-2: alert() вместо inline-ошибок."""
    html = _read("app/templates/user_orders.html")

    assert "alert(" not in html, "ошибки загрузки снова через alert()"
    assert "${file.name}" not in html, "имя файла снова подставляется в разметку"
    assert "preview.innerHTML" not in html
    assert "textContent = file.name" in html or "caption.textContent" in html
    assert ".upload-error" in html, "нет inline-области для ошибок"


def test_service_worker_does_not_cache_html():
    sw = _read("app/static/sw.js")

    assert "networkFirst" not in sw
    assert "cacheFirst" not in sw
    assert "networkOnly" in sw
    # Единственное место, где что-то пишется в кэш — staleWhileRevalidate
    # (он вызывается только для /static/)
    assert sw.count("cache.put(") == 1
    assert "url.pathname.startsWith('/static/')" in sw


def test_export_excel_does_not_leak_token(admin_client, admin_user, sample_order, regular_user):
    resp = admin_client.get("/admin/export_excel")
    assert resp.status_code == 200

    worksheet = load_workbook(io.BytesIO(resp.data)).active
    rows = list(worksheet.iter_rows(values_only=True))

    headers = [cell for cell in rows[0] if cell]
    assert "Токен" not in headers

    tokens = {admin_user.token, regular_user.token}
    for row in rows[1:]:
        values = {cell for cell in row if isinstance(cell, str)}
        assert not (values & tokens), f"токен попал в выгрузку: {values & tokens}"


def test_export_excel_includes_courier_delivery_and_complaints(
    admin_client, db_session, sample_order
):
    from app.models import Complaint

    sample_order.courier_fio = "Курьер Экспорт"
    sample_order.delivered_at = "01.10.2026"
    db_session.add(Complaint(order_id=sample_order.id, text="жалоба-для-выгрузки", status="new"))
    db_session.commit()

    resp = admin_client.get("/admin/export_excel")
    assert resp.status_code == 200

    worksheet = load_workbook(io.BytesIO(resp.data)).active
    rows = list(worksheet.iter_rows(values_only=True))
    headers = list(rows[0])
    assert "Курьер" in headers
    assert "Доставлен" in headers
    assert "Статус жалобы" in headers
    assert "Кол-во жалоб" in headers

    order_row = next(r for r in rows[1:] if r[0] == sample_order.id)
    data = dict(zip(headers, order_row, strict=False))
    assert data["Курьер"] == "Курьер Экспорт"
    assert data["Доставлен"] == "01.10.2026"
    assert data["Кол-во жалоб"] >= 1


def test_export_excel_empty_result_is_valid_headers_only(admin_client):
    resp = admin_client.get("/admin/export_excel", query_string={"search": "zz-empty-export-9x"})
    assert resp.status_code == 200
    assert resp.mimetype.endswith("sheetml.sheet")

    worksheet = load_workbook(io.BytesIO(resp.data)).active
    rows = list(worksheet.iter_rows(values_only=True))
    assert len(rows) == 1, "в пустой выборке — только строка заголовков"
    assert rows[0][0] == "ID"


# ---- Этап 4: mobile-first (нижняя навигация, прогресс загрузки) ----


def test_bottom_nav_present_for_admin_and_absent_for_anonymous(admin_client, app):
    resp = admin_client.get("/admin/")
    assert resp.status_code == 200
    html = resp.get_data(as_text=True)
    assert 'class="bottom-nav"' in html, "нижняя навигация не рендерится админу"
    assert 'class="has-bottom-nav"' in html, "body без отступа под нижнюю навигацию"
    assert "/admin/complaints" in html, "нет пункта «Жалобы»"
    assert "/admin/couriers" in html, "нет пункта «Рейтинг»"
    # Навигация — только существующие маршруты (без javascript:)
    assert "javascript:" not in html

    # Отдельный клиент: тот же test_client хранит сессию admin_client
    anon = app.test_client().get("/login")
    assert 'class="bottom-nav"' not in anon.get_data(as_text=True), "анонимам навигация не нужна"


def test_bottom_nav_courier_links_and_picker_branch(auth_client):
    """Курьер: Главная/Заказы/Чат; нет административных маршрутов."""
    resp = auth_client.get("/orders")
    assert resp.status_code == 200
    html = resp.get_data(as_text=True)
    assert 'class="bottom-nav"' in html
    assert 'href="/dashboard"' in html
    assert 'href="/orders"' in html
    assert 'href="/chat"' in html
    for forbidden in ("/admin", "complaint_filter", "/admin/couriers"):
        assert forbidden not in html, f"курьеру утекает {forbidden}"


def test_bottom_nav_css_fixed_safe_area_and_desktop_hidden():
    css = _read("app/static/style.css")

    assert ".bottom-nav" in css, "нет стилей нижней навигации"
    assert "env(safe-area-inset-bottom" in css, "нет поддержки safe-area (iPhone)"
    assert "@media (max-width: 768px)" in css
    # Панель не перекрывает контент: body получает padding-bottom
    assert "body.has-bottom-nav" in css
    assert "padding-bottom: calc(" in css
    # Тач-цели >= 44px (WCAG 2.5.5)
    assert "min-height: 56px" in css


def test_upload_progress_is_enhancement_with_native_fallback():
    """Прогресс загрузки — XHR enhancement; без JS форма нативная (action/method)."""
    html = _read("app/templates/user_orders.html")

    assert "upload-progress" in html, "нет индикатора прогресса загрузки"
    assert "XMLHttpRequest" in html, "прогресс не реализован через XHR"
    assert "window.location.reload()" in html, "после AJAX нет reload с flash"
    # Нативный fallback сохранён
    assert 'enctype="multipart/form-data"' in html
    assert 'method="post"' in html


def test_long_strings_use_overflow_wrap():
    """Длинные ФИО/номера/жалобы не должны рвать вёрстку по горизонтали."""
    css = _read("app/static/style.css")

    assert "overflow-wrap: anywhere" in css
    assert ".order-description" in css
    assert ".table td" in css


def test_dashboard_attention_metrics_carry_filters(admin_client, db_session, regular_user):
    """§4: карточки проблем дашборда ведут на списки с применёнными фильтрами."""
    from app.models import Order

    db_session.add(
        Order(user_id=regular_user.id, order_number="DASH-1", complaint_text="Принесли чужой заказ")
    )
    db_session.add(
        Order(
            user_id=regular_user.id,
            order_number="DASH-2",
            complaint_text="Не донесли часть товаров из заказа",
        )
    )
    db_session.commit()

    html = admin_client.get("/admin/").get_data(as_text=True)
    assert "Требует внимания" in html
    assert "photo_filter=without_photos" in html
    assert "complaint_filter=with_complaints" in html
    assert "photo_filter=with_photos" in html
    # KPI-карточки — ссылки, а не декоративные div
    assert "kpi-card" in html


def test_dashboard_has_no_placeholder_metrics(admin_client):
    """§4: только реальные данные из admin_stats, без фиктивных цифр."""
    html = admin_client.get("/admin/").get_data(as_text=True)

    assert "demo" not in html.lower()
    assert "example.com" not in html


# ---- Этап 7–8: mobile-first заказы курьера и чат ----


def test_user_orders_template_has_address_and_pending_chip():
    """Этап 7: адрес, подсказка фото-подтемы, чип pending-фильтра."""
    html = _read("app/templates/user_orders.html")

    assert "order-address" in html, "нет блока адреса доставки"
    assert "photo-required-hint" in html, "нет подсказки «нужно фото по жалобе»"
    assert "photo='pending'" in html, "нет чипа-фильтра «Ждут фото»"


def test_chat_template_is_mobile_first():
    """Этап 8: пузыри сообщений и тач-цели вместо activity-list."""
    html = _read("app/templates/chat.html")
    css = _read("app/static/style.css")

    assert "chat-bubble" in html, "нет пузырей сообщений"
    assert "chat-list" in html, "нет контейнера чата"
    assert "activity-item" not in html, "чат снова использует activity-list"
    # iOS-zoom: textarea >= 16px
    assert ".chat-compose textarea" in css
    assert "font-size: 16px" in css
    # Тач-цели
    assert ".chat-publish-btn" in css
    assert "min-height: 44px" in css


def test_courier_dashboard_metrics_carry_filters(auth_client, db_session, regular_user):
    """Этап 9: показатели дашборда курьера ведут на списки с фильтрами (§4)."""
    from app.models import Order

    db_session.add(
        Order(
            user_id=regular_user.id,
            order_number="CDASH-1",
            complaint_text="Принесли чужой заказ",
        )
    )
    db_session.commit()

    html = auth_client.get("/dashboard").get_data(as_text=True)
    # «Ожидает фото» и hero-кнопка → список с фильтром pending
    assert "photo=pending" in html
    assert "kpi-card" in html
    # Дубли шапки скрыты классом (CSS ≤768px)
    assert "dashboard-header-actions" in html
    assert "dashboard-greeting" in html, "нет переноса длинного ФИО"


def test_picker_dashboard_is_mobile_first():
    """Этап 10: дашборд сборщика — перенос ФИО, скрытие дублей шапки, без двойного overflow."""
    html = _read("app/templates/user_picker_dashboard.html")

    assert "dashboard-greeting" in html
    assert "dashboard-header-actions" in html
    assert 'style="overflow-x: auto;"' not in html, "лишний overflow-x ломает карточный вид таблицы"


def test_order_photos_gallery_is_mobile_first():
    """Этап 11: галерея фото — плотная сетка, тач-цели, адрес заказа."""
    html = _read("app/templates/user_order_photos.html")

    assert "photos-grid" in html
    assert "minmax(140px, 1fr)" in html, "на телефонах сетка недостаточно плотная"
    assert "order.address" in html, "нет адреса заказа в блоке информации"
    # Тач-цели lightbox-закрытия и удаления
    assert "width: 44px" in html
    assert "min-height: 44px" in html


def test_mobile_inputs_prevent_ios_zoom():
    """Этап 12: на мобильных все поля ввода >= 16px (иначе iOS автозум при фокусе)."""
    css = _read("app/static/style.css")

    assert "font-size: 16px" in css
    # Правило живёт в media query ≤768px и покрывает select/textarea/input
    mobile_block = css[css.index("max-width: 768px") :]
    assert 'input:not([type="hidden"])' in mobile_block
    assert "select," in mobile_block
    assert "textarea {" in mobile_block
