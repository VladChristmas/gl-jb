"""Регрессионные тесты аудита: карточка жалобы, tz-парсинг дат операционного анализа."""

import io
from datetime import datetime, timedelta, timezone

import openpyxl

from app.models.complaint import Complaint
from app.models.operational_analysis import AssemblyOrder
from app.services.operational_analysis_service import OperationalAnalysisService


def _make_complaint(db_session, sample_order):
    complaint = Complaint(order_id=sample_order.id, text="Повреждён заказ", status="открыта")
    db_session.add(complaint)
    db_session.commit()
    return complaint


def test_complaint_detail_page_renders(admin_client, db_session, sample_order):
    """Страница жалобы должна открываться (200), а не падать 500 из-за несуществующего url_for."""
    complaint = _make_complaint(db_session, sample_order)
    resp = admin_client.get(f"/admin/complaint/{complaint.id}")
    assert resp.status_code == 200, f"expected 200, got {resp.status_code}"


def test_complaint_detail_status_change_post(admin_client, db_session, sample_order):
    """POST-форма смены статуса должна работать (до фикса — 405: роут был только GET)."""
    complaint = _make_complaint(db_session, sample_order)
    resp = admin_client.post(
        f"/admin/complaint/{complaint.id}",
        data={"action": "change_status", "status": "в работе"},
        follow_redirects=False,
    )
    assert resp.status_code == 302, f"expected 302 redirect, got {resp.status_code}"
    db_session.refresh(complaint)
    assert complaint.status == "в работе"


class TestDatetimeTimezoneNormalization:
    """tz-aware даты из ISO-строк должны нормализоваться к naive UTC.

    До фикса datetime.fromisoformat("...Z") возвращал tz-aware значение, и
    `start - queue_time` (naive из Excel) падал с TypeError, роняя весь анализ.
    """

    def test_iso_z_is_naive(self):
        dt = OperationalAnalysisService._parse_datetime("2026-05-01T10:00:00Z")
        assert dt is not None
        assert dt.tzinfo is None
        assert dt == datetime(2026, 5, 1, 10, 0, 0)

    def test_offset_string_converted_to_utc(self):
        dt = OperationalAnalysisService._parse_datetime("2026-05-01T12:00:00+03:00")
        assert dt is not None
        assert dt.tzinfo is None
        assert dt == datetime(2026, 5, 1, 9, 0, 0)

    def test_aware_datetime_input_normalized(self):
        aware = datetime(2026, 5, 1, 10, 0, tzinfo=timezone(timedelta(hours=3)))
        dt = OperationalAnalysisService._parse_datetime(aware)
        assert dt is not None
        assert dt.tzinfo is None
        assert dt == datetime(2026, 5, 1, 7, 0, 0)

    def test_mixed_subtraction_does_not_raise(self):
        start = OperationalAnalysisService._parse_datetime("2026-10-10T09:05:00Z")
        queue_naive = datetime(2026, 10, 10, 9, 0)  # как приходит из Excel-ячейки
        delta = start - queue_naive
        assert delta.total_seconds() / 60.0 == 5.0


def _make_tz_mixed_file() -> io.BytesIO:
    """Файл, где queue_time — naive datetime, а время сборки — ISO-строка с 'Z'.

    До фикса такой файл ронял analyze_file с TypeError.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet(title="Жалобы")
    ws.append(["Номер заказа", "Жалоба"])
    ws.append(["001", "Повреждена упаковка"])

    ws = wb.create_sheet(title="Заказы")
    ws.append(["Id заказа", "Время поступления", "Курьер", "Адрес"])
    ws.append(["001", datetime(2026, 10, 10, 9, 0), "Иванов И.И.", "ул. Ленина, 1"])

    ws = wb.create_sheet(title="Собранные заказы")
    ws.append(["ID заказа", "Сборщик", "Время начала", "Время окончания", "Количество товаров"])
    ws.append(["001", "Козлов А.А.", "2026-10-10T09:05:00Z", "2026-10-10T09:15:00Z", 5])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def test_full_analysis_with_tz_mixed_timestamps(db_session):
    """Полный анализ файла со смешанными tz-датами должен завершиться без ошибки."""
    session = OperationalAnalysisService.analyze_file(_make_tz_mixed_file())

    assert session.status == "completed"
    orders = AssemblyOrder.query.filter_by(session_id=session.id).all()
    assert len(orders) == 1
    assert orders[0].assembly_duration_minutes == 10.0
    assert orders[0].wait_time_minutes == 5.0
