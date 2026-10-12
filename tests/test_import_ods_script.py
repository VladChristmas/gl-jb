"""Тесты scripts/import_ods.py: конвертация ODS, обработка ошибок, полный прогон."""

import io
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from app.services.operational_analysis_service import OperationalAnalysisService
from scripts import import_ods


def _write_sample_ods(path: Path) -> Path:
    complaints = pd.DataFrame(
        {"Номер заказа": ["001", "002"], "Жалоба": ["Повреждена упаковка", "Поздняя доставка"]}
    )
    orders = pd.DataFrame(
        {
            "Id заказа": ["001", "002"],
            "Время поступления": [datetime(2026, 10, 10, 9, 0), datetime(2026, 10, 10, 9, 10)],
            "Курьер": ["Иванов И.И.", "Петров П.П."],
            "Адрес": ["ул. Ленина, 1", None],
        }
    )
    assembly = pd.DataFrame(
        {
            "ID заказа": ["001", "002"],
            "Сборщик": ["Козлов А.А.", "Смирнов Б.Б."],
            "Время начала": [datetime(2026, 10, 10, 9, 5), datetime(2026, 10, 10, 9, 25)],
            "Время окончания": [datetime(2026, 10, 10, 9, 15), datetime(2026, 10, 10, 9, 38)],
            "Количество товаров": [5, 8],
        }
    )
    with pd.ExcelWriter(path, engine="odf") as writer:
        complaints.to_excel(writer, sheet_name="Жалобы", index=False)
        orders.to_excel(writer, sheet_name="Заказы", index=False)
        assembly.to_excel(writer, sheet_name="Собранные заказы", index=False)
    return path


def test_ods_to_xlsx_preserves_sheets_headers_and_types(tmp_path):
    ods = _write_sample_ods(tmp_path / "sample.ods")

    xlsx_bytes = import_ods.ods_to_xlsx_bytes(str(ods))
    assert len(xlsx_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert wb.sheetnames == ["Жалобы", "Заказы", "Собранные заказы"]

    ws = wb["Заказы"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert headers == ["Id заказа", "Время поступления", "Курьер", "Адрес"]

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    # Ведущие нули обязаны сохраняться (иначе связывание между листами рассыпается)
    assert rows[0][0] == "001"
    # Дата после dtype=str становится строкой, но обязана парситься сервисом в тот же момент
    assert OperationalAnalysisService._parse_datetime(rows[0][1]) == datetime(2026, 10, 10, 9, 0)
    assert rows[1][3] is None  # пустая ячейка -> None

    # Числа (количество товаров) также приходят строками и парсятся сервисом
    ws_asm = wb["Собранные заказы"]
    asm_rows = list(ws_asm.iter_rows(min_row=2, values_only=True))
    assert OperationalAnalysisService._parse_number(asm_rows[0][4]) == 5.0


def test_empty_ods_sheets_conversion(tmp_path):
    path = tmp_path / "empty.ods"
    with pd.ExcelWriter(path, engine="odf") as writer:
        pd.DataFrame({"col": []}).to_excel(writer, sheet_name="Жалобы", index=False)

    xlsx_bytes = import_ods.ods_to_xlsx_bytes(str(path))
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert wb.sheetnames == ["Жалобы"]


def test_main_missing_file_exits_1(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(sys, "argv", ["import_ods.py", str(tmp_path / "nope.ods")])
    with pytest.raises(SystemExit) as exc:
        import_ods.main()
    assert exc.value.code == 1
    assert "Файл не найден" in capsys.readouterr().out


def test_main_corrupt_file_exits_1_with_friendly_message(tmp_path, monkeypatch, capsys):
    broken = tmp_path / "broken.ods"
    broken.write_text("this is not an ODS file", encoding="utf-8")

    monkeypatch.setattr(sys, "argv", ["import_ods.py", str(broken)])
    with pytest.raises(SystemExit) as exc:
        import_ods.main()
    assert exc.value.code == 1
    out = capsys.readouterr().out
    assert "Не удалось прочитать ODS-файл" in out
    assert "Traceback" not in out


def test_main_end_to_end_completes(tmp_path, monkeypatch, capsys):
    """Полный прогон: ODS -> XLSX -> анализ -> завершённая сессия, без SystemExit."""
    ods = _write_sample_ods(tmp_path / "e2e.ods")
    monkeypatch.setattr(sys, "argv", ["import_ods.py", str(ods)])

    import_ods.main()

    out = capsys.readouterr().out
    assert "Статус: completed" in out
    assert "Готово. Результаты" in out
