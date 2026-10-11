"""Regression: несколько жалоб на один заказ (P1-6).

Раньше `Order.complaint_text` склеивался через "; ", а
`complaint_status`/`complaint_date` перезаписывались каждой строкой импорта —
статус и дата предыдущих жалоб терялись безвозвратно.
"""

import io
from datetime import UTC, datetime

import openpyxl
from flask import Flask
from flask_migrate import Migrate, downgrade, upgrade

from app.extensions import db
from app.models import Complaint, ImportMapping, Order, User
from app.services.complaint_service import ComplaintService
from app.services.import_service import ImportService


def make_file(headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _app_for(uri: str) -> Flask:
    application = Flask("complaint-migration-test")
    application.config["SQLALCHEMY_DATABASE_URI"] = uri
    db.init_app(application)
    Migrate(application, db)
    return application


class TestComplaintModel:
    def test_three_complaints_on_one_order(self, db_session, regular_user):
        order = Order(user_id=regular_user.id, order_number="C-1")
        db_session.add(order)
        db_session.commit()

        d1 = datetime(2026, 1, 1, 12, 0)
        d2 = datetime(2026, 1, 2, 12, 0)
        d3 = datetime(2026, 1, 3, 12, 0)
        ComplaintService.add(order, "Не вручен вовремя", "открыта", d1)
        ComplaintService.add(order, "Повреждена упаковка", "в работе", d2)
        ComplaintService.add(order, "Не тот состав", "закрыта", d3)
        db_session.commit()

        rows = Complaint.query.filter_by(order_id=order.id).order_by(Complaint.id).all()
        assert [r.text for r in rows] == [
            "Не вручен вовремя",
            "Повреждена упаковка",
            "Не тот состав",
        ]
        assert [r.status for r in rows] == ["открыта", "в работе", "закрыта"]
        assert [r.created_at for r in rows] == [d1, d2, d3]

        db_session.refresh(order)
        assert order.complaint_text == "Не вручен вовремя; Повреждена упаковка; Не тот состав"
        assert order.complaint_status == "закрыта"
        assert order.complaint_date == d3

    def test_duplicate_text_is_not_duplicated(self, db_session, regular_user):
        order = Order(user_id=regular_user.id, order_number="C-2")
        db_session.add(order)
        db_session.commit()

        ComplaintService.add(order, "Одна и та же", "открыта")
        ComplaintService.add(order, "Одна и та же", "закрыта")
        db_session.commit()

        rows = Complaint.query.filter_by(order_id=order.id).all()
        assert len(rows) == 1
        assert rows[0].status == "закрыта"

    def test_legacy_mirror_text_is_not_lost(self, db_session, regular_user):
        """Заказ, у которого жалобы существуют только в зеркале."""
        legacy_date = datetime(2026, 2, 2, 10, 0)
        order = Order(
            user_id=regular_user.id,
            order_number="C-3",
            complaint_text="Старая жалоба",
            complaint_status="закрыта",
            complaint_date=legacy_date,
        )
        db_session.add(order)
        db_session.commit()

        ComplaintService.add(order, "Новая жалоба", "открыта", datetime(2026, 3, 3, 10, 0))
        db_session.commit()

        rows = Complaint.query.filter_by(order_id=order.id).order_by(Complaint.id).all()
        assert [r.text for r in rows] == ["Старая жалоба", "Новая жалоба"]

        db_session.refresh(order)
        assert order.complaint_text == "Старая жалоба; Новая жалоба"
        assert order.complaint_status == "открыта"


class TestComplaintImport:
    def test_import_keeps_status_and_date_of_every_complaint(self, db_session, regular_user):
        order = Order(user_id=regular_user.id, order_number="ORD-900")
        db_session.add(order)
        db_session.commit()

        mapping = ImportMapping(
            name="complaints-multi",
            type="complaints",
            id_column="ID заказа",
            field_mapping={
                "Жалоба": "complaint_text",
                "Статус": "complaint_status",
                "Дата": "complaint_date",
            },
            skip_first_row=True,
        )
        db_session.add(mapping)
        db_session.commit()

        file = make_file(
            ["ID заказа", "Жалоба", "Статус", "Дата"],
            [
                ["ORD-900", "Жалоба 1", "открыта", "2026-01-01"],
                ["ORD-900", "Жалоба 2", "в работе", "2026-01-02"],
                ["ORD-900", "Жалоба 3", "закрыта", "2026-01-03"],
            ],
        )
        result = ImportService.import_with_mapping(file, mapping)

        db_session.refresh(order)
        rows = Complaint.query.filter_by(order_id=order.id).order_by(Complaint.id).all()
        assert [r.text for r in rows] == ["Жалоба 1", "Жалоба 2", "Жалоба 3"]
        assert [r.status for r in rows] == ["открыта", "в работе", "закрыта"]
        assert [r.created_at for r in rows] == [
            datetime(2026, 1, 1),
            datetime(2026, 1, 2),
            datetime(2026, 1, 3),
        ]

        assert order.complaint_text == "Жалоба 1; Жалоба 2; Жалоба 3"
        assert order.complaint_status == "закрыта"
        assert order.complaint_date == datetime(2026, 1, 3)
        assert result["skipped"] == 0
        assert result["updated"] == 3


def test_migration_backfills_legacy_mirror_complaints(tmp_path):
    """Миграция переносит склеенный `complaint_text` в строки `complaints`."""
    uri = "sqlite:///" + (tmp_path / "legacy-complaints.db").as_posix()
    application = _app_for(uri)

    legacy_date = datetime(2026, 2, 2, 10, 0)

    with application.app_context():
        upgrade()

        user = User(fio="Легаси", token="legacy-complaints-token")
        db.session.add(user)
        db.session.flush()
        order = Order(
            user_id=user.id,
            order_number="L-1",
            complaint_text="Жалоба A; Жалоба B; Жалоба C",
            complaint_status="закрыта",
            complaint_date=legacy_date,
        )
        db.session.add(order)
        db.session.commit()
        order_id = order.id

        # Симулируем повторный прогон миграции жалоб на легаси-данных
        downgrade(revision="a1b2c3d4e5f6")
        upgrade()

        rows = Complaint.query.filter_by(order_id=order_id).order_by(Complaint.id).all()
        assert [r.text for r in rows] == ["Жалоба A", "Жалоба B", "Жалоба C"]
        assert [r.status for r in rows] == ["закрыта", "закрыта", "закрыта"]
        assert [r.created_at for r in rows] == [legacy_date, legacy_date, legacy_date]

        order = db.session.get(Order, order_id)
        assert order.complaint_text == "Жалоба A; Жалоба B; Жалоба C"
        assert order.complaint_status == "закрыта"


class TestComplaintsPage:
    """Этап 6: /admin/complaints — только чтение, фильтры, ссылки на заказ."""

    def test_requires_admin(self, app):
        anon = app.test_client()
        resp = anon.get("/admin/complaints")
        assert resp.status_code == 302
        assert "/admin/login" in resp.headers["Location"]

    def test_lists_complaints_with_order_link(self, admin_client, db_session, regular_user):
        order = Order(user_id=regular_user.id, order_number="CP-1")
        db_session.add(order)
        db_session.commit()
        ComplaintService.add(order, "Товар побит/вскрыт", "открыта", datetime(2026, 3, 1, 10, 0))
        ComplaintService.add(order, "Не донесли часть товаров", "закрыта")
        db_session.commit()

        html = admin_client.get("/admin/complaints").get_data(as_text=True)
        assert "Товар побит/вскрыт" in html
        assert "Не донесли часть товаров" in html
        assert f"/admin/orders/edit/{order.id}" in html
        assert "открыта" in html
        assert "закрыта" in html
        # Фото-подтема без фото → честный бейдж «ждём фото»
        assert "ждём фото" in html

    def test_status_filter(self, admin_client, db_session, regular_user):
        order = Order(user_id=regular_user.id, order_number="CP-2")
        db_session.add(order)
        db_session.commit()
        ComplaintService.add(order, "Принесли чужой заказ", "открыта")
        ComplaintService.add(order, "Не учли комментарий", "закрыта")
        db_session.commit()

        html = admin_client.get("/admin/complaints?status=открыта").get_data(as_text=True)
        assert "Принесли чужой заказ" in html
        assert "Не учли комментарий" not in html

        # Неизвестный статус → фильтр игнорируется, 200 (без 500)
        resp = admin_client.get("/admin/complaints?status=хакерский")
        assert resp.status_code == 200
        assert "Принесли чужой заказ" in resp.get_data(as_text=True)

    def test_no_status_filter(self, admin_client, db_session, regular_user):
        order = Order(user_id=regular_user.id, order_number="CP-3")
        db_session.add(order)
        db_session.commit()
        ComplaintService.add(order, "Жалоба без статуса", None)
        ComplaintService.add(order, "Жалоба закрытая", "закрыта")
        db_session.commit()

        html = admin_client.get("/admin/complaints?status=none").get_data(as_text=True)
        assert "Жалоба без статуса" in html
        assert "Жалоба закрытая" not in html

    def test_search_by_order_number(self, admin_client, db_session, regular_user):
        order = Order(user_id=regular_user.id, order_number="FIND-ME-42")
        db_session.add(order)
        db_session.commit()
        ComplaintService.add(order, "Товар побит/вскрыт", "открыта")
        db_session.commit()

        html = admin_client.get("/admin/complaints?search=FIND-ME-42").get_data(as_text=True)
        assert "Товар побит/вскрыт" in html

        html = admin_client.get("/admin/complaints?search=NoSuchOrder").get_data(as_text=True)
        assert "Жалоб нет" in html

    def test_courier_filter(self, admin_client, db_session, regular_user):
        order = Order(
            user_id=regular_user.id,
            order_number="CP-4",
            courier_fio="Курьер Фильтр",
            # Дата обязательна: заказы с курьером без даты попадают в
            # orders_without_date рейтинга (session-scoped БД в тестах)
            delivered_at=datetime(2026, 3, 2, 12, 0),
        )
        db_session.add(order)
        db_session.commit()
        ComplaintService.add(order, "Товар побит/вскрыт", "открыта")
        db_session.commit()

        html = admin_client.get("/admin/complaints?courier=Курьер Фильтр").get_data(as_text=True)
        assert "Товар побит/вскрыт" in html

        html = admin_client.get("/admin/complaints?courier=Несуществующий").get_data(as_text=True)
        assert "Жалоб нет" in html

    def test_status_choices_from_data(self, db_session, regular_user):
        order = Order(user_id=regular_user.id, order_number="CP-5")
        db_session.add(order)
        db_session.commit()
        ComplaintService.add(order, "A", "открыта")
        ComplaintService.add(order, "B", "закрыта")
        ComplaintService.add(order, "C", None)
        db_session.commit()

        choices = ComplaintService.status_choices()
        values = [value for value, _label, _count in choices]
        assert "открыта" in values
        assert "закрыта" in values
        assert ComplaintService.NO_STATUS in values
        labels = {value: label for value, label, _count in choices}
        assert labels["открыта"] == "Открыта"
        assert labels[ComplaintService.NO_STATUS] == "Без статуса"

    def test_mirror_only_count(self, db_session, regular_user):
        # Session-scoped БД: считаем прирост, а не абсолют (хвосты других тестов)
        before = ComplaintService.mirror_only_count()
        legacy = Order(
            user_id=regular_user.id,
            order_number="CP-legacy",
            complaint_text="Только зеркало",
        )
        db_session.add(legacy)
        hydrated = Order(user_id=regular_user.id, order_number="CP-hydrated")
        db_session.add(hydrated)
        db_session.commit()
        ComplaintService.add(hydrated, "Строка complaints", "открыта")
        db_session.commit()

        assert ComplaintService.mirror_only_count() == before + 1

    def test_page_is_readonly_no_status_mutation(self, admin_client, db_session, regular_user):
        """§ Правила: статусы жалоб не меняются никаким GET/страницей."""
        order = Order(user_id=regular_user.id, order_number="CP-6")
        db_session.add(order)
        db_session.commit()
        complaint = ComplaintService.add(order, "Товар побит/вскрыт", "открыта")
        db_session.commit()
        assert complaint is not None
        cid = complaint.id

        admin_client.get("/admin/complaints")
        admin_client.get("/admin/complaints?status=открыта")
        admin_client.get("/admin/complaints?search=Товар")

        db_session.expire_all()
        row = db_session.get(Complaint, cid)
        assert row is not None
        assert row.status == "открыта"

    def test_date_filter_by_created_at(self, admin_client, db_session, regular_user):
        """Этап 13: date_from/date_to фильтруют по дате создания жалобы."""
        order = Order(user_id=regular_user.id, order_number="CP-7")
        db_session.add(order)
        db_session.commit()
        ComplaintService.add(order, "старая жалоба", "открыта", datetime(2026, 1, 10, 10, 0))
        ComplaintService.add(order, "свежая жалоба", "открыта", datetime(2026, 3, 15, 10, 0))
        db_session.commit()

        html = admin_client.get(
            "/admin/complaints?date_from=2026-03-01&date_to=2026-03-31"
        ).get_data(as_text=True)
        assert "свежая жалоба" in html
        assert "старая жалоба" not in html

        html = admin_client.get("/admin/complaints?date_to=2026-02-01").get_data(as_text=True)
        assert "старая жалоба" in html
        assert "свежая жалоба" not in html

    def test_period_preset_filters_by_day(self, admin_client, db_session, regular_user):
        """Этап 13: period=day показывает только жалобы за сегодня."""
        from datetime import datetime as dt

        order = Order(user_id=regular_user.id, order_number="CP-8")
        db_session.add(order)
        db_session.commit()
        ComplaintService.add(order, "жалоба сегодня", "открыта", dt.now(UTC).replace(tzinfo=None))
        ComplaintService.add(order, "жалоба вчера", "открыта", dt(2020, 1, 1, 10, 0))
        db_session.commit()

        html = admin_client.get("/admin/complaints?period=day").get_data(as_text=True)
        assert "жалоба сегодня" in html
        assert "жалоба вчера" not in html

    def test_invalid_period_ignored(self, admin_client, db_session, regular_user):
        """Этап 13: мусорный period не ломает страницу и не фильтрует."""
        order = Order(user_id=regular_user.id, order_number="CP-9")
        db_session.add(order)
        db_session.commit()
        ComplaintService.add(order, "жалоба x", "открыта")
        db_session.commit()

        resp = admin_client.get("/admin/complaints?period=__hack__")
        assert resp.status_code == 200
        assert "жалоба x" in resp.get_data(as_text=True)


class TestComplaintsExport:
    """Этап 14: экспорт текущей выборки жалоб в Excel (§12)."""

    def test_export_follows_current_filter(self, admin_client, db_session, regular_user):
        order = Order(user_id=regular_user.id, order_number="EX-1")
        db_session.add(order)
        db_session.commit()
        ComplaintService.add(order, "жалоба для экспорта", "открыта", datetime(2026, 5, 1, 10, 0))
        ComplaintService.add(order, "другая жалоба", "закрыта", datetime(2020, 1, 1, 10, 0))
        db_session.commit()

        resp = admin_client.get("/admin/complaints/export_excel?date_from=2026-05-01")
        assert resp.status_code == 200
        assert (
            resp.content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "attachment" in resp.headers.get("Content-Disposition", "")

        wb = openpyxl.load_workbook(io.BytesIO(resp.data))
        ws = wb.active
        assert ws is not None
        assert ws.title == "Жалобы"
        rows = list(ws.iter_rows(values_only=True))
        texts = [r[3] for r in rows[1:]]  # колонка «Текст жалобы»
        assert "жалоба для экспорта" in texts
        assert "другая жалоба" not in texts

    def test_export_all_when_no_filter(self, admin_client, db_session, regular_user):
        order = Order(user_id=regular_user.id, order_number="EX-2")
        db_session.add(order)
        db_session.commit()
        ComplaintService.add(order, "жалоба А", "открыта")
        ComplaintService.add(order, "жалоба Б", "в работе")
        db_session.commit()

        resp = admin_client.get("/admin/complaints/export_excel")
        assert resp.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(resp.data))
        ws = wb.active
        assert ws is not None
        rows = list(ws.iter_rows(values_only=True))
        texts = {r[3] for r in rows[1:]}
        assert {"жалоба А", "жалоба Б"} <= texts

    def test_export_requires_admin(self, app):
        anon = app.test_client()
        resp = anon.get("/admin/complaints/export_excel")
        assert resp.status_code == 302
        assert "/admin/login" in resp.headers["Location"]


class TestCourierSummary:
    """Этап 15: топ курьеров по жалобам (только отображение, рейтинг не трогаем)."""

    def test_summary_groups_by_courier_and_links_to_filter(
        self, admin_client, db_session, regular_user
    ):
        # Дата обязательна: заказы с курьером без даты портят KPI orders_without_date
        order1 = Order(
            user_id=regular_user.id,
            order_number="CS-1",
            courier_fio="Иванов И.И.",
            delivered_at=datetime(2026, 5, 1, 12, 0),
        )
        order2 = Order(
            user_id=regular_user.id,
            order_number="CS-2",
            courier_fio="Иванов И.И.",
            delivered_at=datetime(2026, 5, 2, 12, 0),
        )
        order3 = Order(
            user_id=regular_user.id,
            order_number="CS-3",
            courier_fio="Петров П.П.",
            delivered_at=datetime(2026, 5, 3, 12, 0),
        )
        db_session.add_all([order1, order2, order3])
        db_session.commit()
        ComplaintService.add(order1, "жалоба 1", "открыта")
        ComplaintService.add(order2, "жалоба 2", "открыта")
        ComplaintService.add(order3, "жалоба 3", "открыта")
        db_session.commit()

        html = admin_client.get("/admin/complaints").get_data(as_text=True)
        assert "Топ курьеров по жалобам" in html
        assert "Иванов И.И." in html
        # Ссылка ведёт на список с фильтром courier (§4)
        assert "courier=%D0%98%D0%B2%D0%B0%D0%BD%D0%BE%D0%B2" in html or "courier=Иванов" in html

    def test_summary_respects_status_filter(self, admin_client, db_session, regular_user):
        order = Order(
            user_id=regular_user.id,
            order_number="CS-4",
            courier_fio="Сидоров С.С.",
            delivered_at=datetime(2026, 5, 4, 12, 0),
        )
        db_session.add(order)
        db_session.commit()
        ComplaintService.add(order, "открытая жалоба", "открыта")
        ComplaintService.add(order, "закрытая жалоба", "закрыта")
        db_session.commit()

        html = admin_client.get("/admin/complaints?status=открыта").get_data(as_text=True)
        # В сводке только жалобы со статусом «открыта»: у Сидорова 1
        assert "Сидоров С.С." in html
        assert html.count("Сидоров С.С.") >= 2  # фильтр-селект + сводка

    def test_summary_respects_date_filter(self, admin_client, db_session, regular_user):
        order = Order(
            user_id=regular_user.id,
            order_number="CS-5",
            courier_fio="Кузнецов К.К.",
            delivered_at=datetime(2026, 6, 1, 12, 0),
        )
        db_session.add(order)
        db_session.commit()
        ComplaintService.add(order, "старая", "открыта", datetime(2020, 1, 1, 10, 0))
        ComplaintService.add(order, "свежая", "открыта", datetime(2026, 6, 1, 10, 0))
        db_session.commit()

        html = admin_client.get("/admin/complaints?date_from=2026-06-01").get_data(as_text=True)
        assert "Кузнецов К.К." in html

    def test_summary_empty_when_no_couriers(self, admin_client, db_session, regular_user):
        order = Order(user_id=regular_user.id, order_number="CS-NO-1")  # без курьера
        db_session.add(order)
        db_session.commit()
        ComplaintService.add(order, "жалоба без курьера unique-no-courier", "открыта")
        db_session.commit()

        # Изолируем выборку поиском: у неё нет курьера → сводки нет
        html = admin_client.get("/admin/complaints?search=unique-no-courier").get_data(as_text=True)
        assert "жалоба без курьера" in html
        assert "Топ курьеров по жалобам" not in html
