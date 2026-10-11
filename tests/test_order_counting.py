"""Регрессионные тесты подсчёта и отображения заказов (§11 аудита).

12 обязательных сценариев:
1. 100 уникальных заказов отображаются как 100;
2. заказ не считается дважды (фото/жалобы/join);
3. повторный импорт не удваивает заказы;
4. заказ с тремя жалобами считается как один заказ;
5. три жалобы считаются как три жалобы;
6. фильтр по дате (полуинтервал границ дня);
7. фильтр по курьеру;
8. комбинация фильтров;
9. пустой результат = 0 без ошибки;
10. деление на ноль не даёт 500;
11. пагинация не меняет общее количество;
12. экспорт после фильтра выгружает выбранную выборку.
"""

import io
import uuid
from datetime import datetime

import openpyxl
import pytest

from app.extensions import db, limiter
from app.models.order import Order
from app.models.photo import Photo
from app.services.complaint_service import ComplaintService
from app.services.dashboard_service import DashboardService
from app.services.import_service import ImportService
from app.services.order_service import OrderService


@pytest.fixture(autouse=True)
def _reset_limiter():
    limiter.reset()
    yield
    limiter.reset()


def _marker() -> str:
    return "KM" + uuid.uuid4().hex[:10]


def _order(db_session, user, order_number, **kwargs) -> Order:
    order = Order(user_id=user.id, order_number=order_number, **kwargs)
    db_session.add(order)
    db_session.commit()
    return order


def _numbers(db_session, marker: str, **filters):
    pagination, items = OrderService.get_all_orders(search=marker, per_page=500, **filters)
    return pagination, [item["order"].order_number for item in items]


def _courier_row(courier_fio: str) -> dict:
    rows = DashboardService.best_couriers(limit=1000)
    return next(row for row in rows if row["fio"] == courier_fio)


class TestCountingRules:
    def test_100_unique_orders_displayed_as_100(self, db_session, regular_user, admin_client):
        # 1. 100 уникальных заказов отображаются как 100
        marker = _marker()
        try:
            for i in range(100):
                db_session.add(Order(user_id=regular_user.id, order_number=f"{marker}-{i:03d}"))
            db_session.commit()

            pagination, numbers = _numbers(db_session, marker)
            assert pagination.total == 100
            assert len(numbers) == 100
            assert len(set(numbers)) == 100

            html = admin_client.get(f"/admin/orders?search={marker}").get_data(as_text=True)
            assert "Заказов по текущему фильтру:</strong> 100" in html
        finally:
            _cleanup_marker(db_session, marker)

    def test_order_not_counted_twice_with_photos_and_complaints(self, db_session, regular_user):
        # 2. заказ не считается дважды при наличии фото и жалоб
        marker = _marker()
        courier = f"Курьер {marker}"
        order = _order(db_session, regular_user, f"{marker}-1", courier_fio=courier)
        db_session.add(
            Photo(order_id=order.id, filename=f"{marker}.jpg", original_filename=f"{marker}.jpg")
        )
        db_session.commit()
        ComplaintService.add(order, f"жалоба {marker}-1")
        ComplaintService.add(order, f"жалоба {marker}-2")
        db_session.commit()

        pagination, numbers = _numbers(db_session, marker)
        assert pagination.total == 1
        assert numbers == [order.order_number]

        # join фото не плодит строки
        pagination, numbers = _numbers(db_session, marker, photo_filter="with_photos")
        assert pagination.total == 1

        # рейтинг курьера: один заказ, не два
        row = _courier_row(courier)
        assert row["orders_count"] == 1
        assert row["orders_with_complaints_count"] == 1
        assert row["complaints_count"] == 2

    def test_reimport_does_not_double_orders(self, db_session, regular_user):
        # 3. повторный импорт не удваивает заказы
        marker = _marker()
        rows = [[regular_user.fio, f"{marker}-{i}", f"описание {marker}"] for i in range(3)]
        headers = ["ФИО", "Номер заказа", "Описание"]

        def make_file():
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(headers)
            for row in rows:
                ws.append(row)
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            return output

        first = ImportService.import_orders_from_excel(make_file(), fio_col_idx=0)
        assert first["imported"] == 3
        assert first["duplicates"] == 0

        second = ImportService.import_orders_from_excel(make_file(), fio_col_idx=0)
        assert second["imported"] == 0
        assert second["duplicates"] == 3

        pagination, numbers = _numbers(db_session, marker)
        assert pagination.total == 3
        assert len(numbers) == 3

    def test_order_with_three_complaints_counts_as_one_order(self, db_session, regular_user):
        # 4. заказ с тремя жалобами считается как один заказ
        marker = _marker()
        courier = f"Курьер {marker}"
        order = _order(db_session, regular_user, f"{marker}-1", courier_fio=courier)
        for i in range(3):
            ComplaintService.add(order, f"жалоба {marker} {i}")
        db_session.commit()

        row = _courier_row(courier)
        assert row["orders_count"] == 1
        assert row["orders_with_complaints_count"] == 1

        # страница жалоб: заказ встречается один раз
        complaint_orders = DashboardService.complaint_couriers()
        matches = [r for r in complaint_orders if r["order_id"] == order.id]
        assert len(matches) == 1

    def test_three_complaints_count_as_three_complaints(self, db_session, regular_user):
        # 5. три жалобы считаются как три жалобы
        marker = _marker()
        courier = f"Курьер {marker}"
        order = _order(db_session, regular_user, f"{marker}-1", courier_fio=courier)
        for i in range(3):
            ComplaintService.add(order, f"жалоба {marker} {i}")
        db_session.commit()

        row = _courier_row(courier)
        assert row["complaints_count"] == 3
        assert row["orders_with_complaints_count"] == 1
        assert row["complaint_rate"] == 100.0


class TestFilters:
    def test_date_filter_uses_half_open_interval(self, db_session, regular_user):
        # 6. фильтр по дате: [начало дня, конец дня + 1) — границы не теряются
        marker = _marker()
        _order(
            db_session,
            regular_user,
            f"{marker}-before",
            created_at=datetime(2026, 1, 9, 23, 59, 59),
        )
        _order(
            db_session, regular_user, f"{marker}-start", created_at=datetime(2026, 1, 10, 0, 0, 0)
        )
        _order(
            db_session, regular_user, f"{marker}-end", created_at=datetime(2026, 1, 10, 23, 59, 59)
        )
        _order(
            db_session, regular_user, f"{marker}-after", created_at=datetime(2026, 1, 11, 0, 0, 0)
        )

        _, numbers = _numbers(db_session, marker, date_from="2026-01-10", date_to="2026-01-10")
        assert sorted(numbers) == [f"{marker}-end", f"{marker}-start"]

        _, numbers = _numbers(db_session, marker, date_from="2026-01-10", date_to="2026-01-11")
        assert sorted(numbers) == [f"{marker}-after", f"{marker}-end", f"{marker}-start"]

        # некорректная дата не роняет запрос
        pagination, numbers = _numbers(db_session, marker, date_from="10.01.2026")
        assert pagination.total == 4

    def test_courier_filter(self, db_session, regular_user, admin_client):
        # 7. фильтр по курьеру
        marker = _marker()
        courier_a = f"Курьер А {marker}"
        courier_b = f"Курьер Б {marker}"
        _order(db_session, regular_user, f"{marker}-a", courier_fio=courier_a)
        _order(db_session, regular_user, f"{marker}-b", courier_fio=courier_a)
        _order(db_session, regular_user, f"{marker}-c", courier_fio=courier_b)

        _, numbers = _numbers(db_session, marker, courier=courier_a)
        assert sorted(numbers) == [f"{marker}-a", f"{marker}-b"]

        html = admin_client.get(
            "/admin/orders", query_string={"search": marker, "courier": courier_b}
        ).get_data(as_text=True)
        assert f"{marker}-c" in html
        assert f"{marker}-a" not in html

    def test_combined_filters_intersection(self, db_session, regular_user):
        # 8. комбинация фильтров = пересечение выборок
        marker = _marker()
        courier_a = f"Курьер А {marker}"
        courier_b = f"Курьер Б {marker}"
        keep = _order(
            db_session,
            regular_user,
            f"{marker}-keep",
            courier_fio=courier_a,
            created_at=datetime(2026, 2, 5, 12, 0, 0),
        )
        _order(
            db_session,
            regular_user,
            f"{marker}-nocomplaint",
            courier_fio=courier_a,
            created_at=datetime(2026, 2, 5, 12, 0, 0),
        )
        _order(
            db_session,
            regular_user,
            f"{marker}-wrong-date",
            courier_fio=courier_a,
            created_at=datetime(2026, 3, 5, 12, 0, 0),
        )
        _order(
            db_session,
            regular_user,
            f"{marker}-other-courier",
            courier_fio=courier_b,
            created_at=datetime(2026, 2, 5, 12, 0, 0),
        )
        ComplaintService.add(keep, f"жалоба {marker}")
        ComplaintService.add(keep, f"жалоба {marker} 2")
        db_session.commit()

        _, numbers = _numbers(
            db_session,
            marker,
            courier=courier_a,
            complaint_filter="with_complaints",
            date_from="2026-02-01",
            date_to="2026-02-28",
            photo_filter="",
        )
        assert numbers == [f"{marker}-keep"]

        # пустое пересечение — ноль, не ошибка
        pagination, numbers = _numbers(
            db_session,
            marker,
            courier=courier_b,
            complaint_filter="with_complaints",
            date_from="2026-02-01",
            date_to="2026-02-28",
        )
        assert pagination.total == 0
        assert numbers == []

    def test_empty_result_returns_zero_without_error(self, db_session, admin_client):
        # 9. пустой результат = 0, без 500
        missing = _marker()
        pagination, numbers = _numbers(db_session, missing)
        assert pagination.total == 0
        assert numbers == []

        resp = admin_client.get(f"/admin/orders?search={missing}")
        assert resp.status_code == 200
        html = resp.get_data(as_text=True)
        assert "Заказов по текущему фильтру:</strong> 0" in html
        assert "Заказов нет" in html

    def test_no_division_by_zero_anywhere(self, db_session, regular_user, admin_client):
        # 10. деление на ноль не приводит к 500
        # пустая база и база с данными — везде finite rate
        for entry in DashboardService.best_couriers(limit=1000):
            assert isinstance(entry["complaint_rate"], float)
        marker = _marker()
        courier = f"Курьер {marker}"
        _order(db_session, regular_user, f"{marker}-1", courier_fio=courier)
        db_session.commit()

        row = _courier_row(courier)
        assert isinstance(row["complaint_rate"], float)
        assert 0.0 <= row["complaint_rate"] <= 100.0

        # курьер без единого заказа: completion_pct не делит на ноль
        from app.models.user import User

        empty_user = User(fio=f"Пустой {marker}", token=f"EMP{uuid.uuid4().hex[:10]}")
        db_session.add(empty_user)
        db_session.commit()
        stats = DashboardService.courier_stats(empty_user.id)
        assert stats["total_orders"] == 0
        assert stats["completion_pct"] == 100

        assert admin_client.get("/admin/couriers").status_code == 200
        assert admin_client.get("/admin/").status_code == 200

    def test_pagination_does_not_change_total(self, db_session, regular_user):
        # 11. пагинация не меняет общее количество
        marker = _marker()
        try:
            for i in range(25):
                db_session.add(Order(user_id=regular_user.id, order_number=f"{marker}-{i:02d}"))
            db_session.commit()

            seen = []
            totals = []
            for page in (1, 2, 3):
                pagination, items = OrderService.get_all_orders(
                    search=marker, page=page, per_page=10
                )
                totals.append(pagination.total)
                seen.extend(item["order"].order_number for item in items)

            assert totals == [25, 25, 25]
            assert len(seen) == 25
            assert len(set(seen)) == 25
        finally:
            _cleanup_marker(db_session, marker)

    def test_export_matches_filtered_selection(self, db_session, regular_user, admin_client):
        # 12. экспорт после фильтра = выбранная выборка
        marker = _marker()
        courier_a = f"Курьер А {marker}"
        courier_b = f"Курьер Б {marker}"
        try:
            for i in range(5):
                db_session.add(
                    Order(
                        user_id=regular_user.id,
                        order_number=f"{marker}-a{i}",
                        courier_fio=courier_a,
                    )
                )
            for i in range(3):
                db_session.add(
                    Order(
                        user_id=regular_user.id,
                        order_number=f"{marker}-b{i}",
                        courier_fio=courier_b,
                    )
                )
            db_session.commit()

            resp = admin_client.get("/admin/export_excel", query_string={"search": marker})
            assert resp.status_code == 200
            assert _xlsx_data_rows(resp.data) == 8

            resp = admin_client.get(
                "/admin/export_excel",
                query_string={"search": marker, "courier": courier_a},
            )
            assert resp.status_code == 200
            assert _xlsx_data_rows(resp.data) == 5
        finally:
            _cleanup_marker(db_session, marker)


class TestImportDuplicates:
    def test_mapping_import_counts_duplicate_rows(self, db_session, regular_user):
        """§8: дубли строк внутри файла показываются в отчёте импорта."""
        from app.models import ImportMapping

        marker = _marker()
        mapping = ImportMapping(
            name=f"mapping {marker}",
            type="orders",
            id_column="ID",
            fio_column="ФИО",
            order_number_column="Номер заказа",
            field_mapping={},
            skip_first_row=True,
        )
        db_session.add(mapping)
        db_session.commit()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["ID", "ФИО", "Номер заказа"])
        ws.append(["EXT-1", regular_user.fio, f"{marker}-1"])
        ws.append(["EXT-1", regular_user.fio, f"{marker}-1"])
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        result = ImportService.import_with_mapping(output, mapping)
        assert result["imported"] == 1
        assert result["duplicates"] == 1
        assert result["updated"] == 1

        pagination, numbers = _numbers(db_session, marker)
        assert pagination.total == 1

    def test_mapping_import_number_conflict_reported_not_crashed(self, db_session, regular_user):
        """§1: заказ с тем же (пользователь, номер), но другим external_id —
        помечается дублем, а не роняет весь импорт IntegrityError'ом."""
        from app.models import ImportMapping

        marker = _marker()
        _order(db_session, regular_user, f"{marker}-1", external_id="OLD-EXT")
        mapping = ImportMapping(
            name=f"mapping {marker} conflict",
            type="orders",
            id_column="ID",
            fio_column="ФИО",
            order_number_column="Номер заказа",
            field_mapping={},
            skip_first_row=True,
        )
        db_session.add(mapping)
        db_session.commit()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["ID", "ФИО", "Номер заказа"])
        ws.append(["NEW-EXT", regular_user.fio, f"{marker}-1"])
        ws.append(["NEW-EXT-2", regular_user.fio, f"{marker}-2"])
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        result = ImportService.import_with_mapping(output, mapping)
        assert result["duplicates"] == 1
        assert result["imported"] == 1
        assert any("уже существует" in error or "занят" in error for error in result["errors"])

        pagination, numbers = _numbers(db_session, marker)
        assert sorted(numbers) == [f"{marker}-1", f"{marker}-2"]


def _xlsx_data_rows(data: bytes) -> int:
    """Число строк данных (без заголовка) в выгруженном xlsx."""
    workbook = openpyxl.load_workbook(io.BytesIO(data))
    sheet = workbook.active
    return max(sheet.max_row - 1, 0)


def _cleanup_marker(db_session, marker: str) -> None:
    orders = Order.query.filter(Order.order_number.like(f"{marker}%")).all()
    for order in orders:
        db.session.delete(order)
    db_session.commit()
